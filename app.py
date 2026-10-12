"""
Geometry Home Delivery Management System
Multi-company | PostgreSQL (Railway) + SQLite (local) | PDF delivery notes
Built on the same foundation as the GH invoicing system.
"""
import os, io, re, json, time, base64, secrets, zipfile
from datetime import datetime, timedelta, timezone
from functools import wraps
from xml.sax.saxutils import escape as xml_escape
from flask import (Flask, g, render_template, request, redirect, url_for,
                   session, flash, jsonify, send_file, send_from_directory)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from db import get_db, last_insert_id, returning_id, USE_PG

BASE_DIR    = os.path.dirname(os.path.abspath(__file__))
# Persistent data folder. On Railway attach a Volume: Railway exposes its mount path as
# RAILWAY_VOLUME_MOUNT_PATH (or set DATA_DIR yourself). Locally this stays inside the project folder.
DATA_DIR    = os.environ.get('DATA_DIR') or os.environ.get('RAILWAY_VOLUME_MOUNT_PATH') or ''
ARCHIVE_DIR = os.path.join(DATA_DIR or BASE_DIR, "Delivery Notes Archive")
UPLOAD_DIR  = os.path.join(DATA_DIR, "uploads") if DATA_DIR else os.path.join(BASE_DIR, "static", "uploads")
BACKUP_DIR  = os.path.join(DATA_DIR or BASE_DIR, "Backups")
IMAGES_DIR  = os.path.join(BASE_DIR, "static", "images")

app = Flask(__name__)
_secret = os.environ.get('SECRET_KEY')
if not _secret:
    import logging
    logging.getLogger(__name__).warning("SECRET_KEY is not set - using a random key; everyone is logged out on every restart. "
                                        "Set SECRET_KEY in the Railway Variables tab.")
    _secret = secrets.token_hex(32)
app.secret_key = _secret
app.config['MAX_CONTENT_LENGTH'] = 512 * 1024 * 1024
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=14)
if os.environ.get('RAILWAY_ENVIRONMENT') or os.environ.get('RAILWAY_PROJECT_ID'):
    app.config['SESSION_COOKIE_SECURE'] = True
    from werkzeug.middleware.proxy_fix import ProxyFix          # Railway terminates HTTPS in front of the app
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
elif os.environ.get('BEHIND_HTTPS_PROXY') == '1':
    # Docker + Cloudflare Tunnel / reverse proxy: the proxy sends X-Forwarded-Proto. The session cookie is marked
    # Secure only on requests that really arrived over HTTPS, so the same container works on plain-http LAN too.
    from werkzeug.middleware.proxy_fix import ProxyFix
    from flask.sessions import SecureCookieSessionInterface
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
    class _AutoSecureSession(SecureCookieSessionInterface):
        def get_cookie_secure(self, app):
            return request.is_secure
    app.session_interface = _AutoSecureSession()

for d in [ARCHIVE_DIR, IMAGES_DIR,
          os.path.join(UPLOAD_DIR, "logos"),
          os.path.join(UPLOAD_DIR, "signatures"),
          os.path.join(UPLOAD_DIR, "stamps")]:
    os.makedirs(d, exist_ok=True)

if DATA_DIR:
    # First start on a fresh volume: copy the bundled logos / signatures / stamps into it (never overwrites).
    import shutil
    _seed = os.path.join(BASE_DIR, "static", "uploads")
    for _root, _dirs, _files in os.walk(_seed):
        _dest = os.path.join(UPLOAD_DIR, os.path.relpath(_root, _seed))
        os.makedirs(_dest, exist_ok=True)
        for _f in _files:
            if not os.path.exists(os.path.join(_dest, _f)):
                shutil.copy2(os.path.join(_root, _f), os.path.join(_dest, _f))

@app.route('/static/uploads/<path:filename>')
def uploaded_file(filename):
    """Serves logos / signatures / stamps from the persistent folder (also works without a volume)."""
    return send_from_directory(UPLOAD_DIR, filename, max_age=3600)

@app.route('/health')
def health():
    return 'ok', 200

ALLOWED = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
def allowed_file(fn):
    return '.' in fn and fn.rsplit('.', 1)[1].lower() in ALLOWED

DEFAULT_DN_NOTES = (
    "Notes: Any discrepancies in the delivery should be reported within 3 days. "
    "This delivery note must be signed and returned to the Supplier.\n"
    "Responsibility for Damage: The driver shall be fully responsible for any damage to the goods during transit, "
    "loading, unloading, or delivery. Any damage resulting from the driver's actions, negligence, or failure to "
    "properly handle the goods will be the driver's sole responsibility, and the driver or their employer may be "
    "held liable for the costs.\n"
    "Acceptance of Goods: The recipient agrees that by signing this delivery note, they have received the goods "
    "listed in good condition and as per the order specifications. Any damages must be reported within 24 hours "
    "of receipt."
)

UNITS = ['PCS', 'SET', 'NOS', 'PAIR', 'BOX', 'CTN', 'ROLL', 'SQM', 'LM', 'KG']


# ── DATABASE INIT ──────────────────────────────────────────────────────────
SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL,
    full_name TEXT, email TEXT, role TEXT DEFAULT 'user',
    is_active INTEGER DEFAULT 1,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP, last_login TEXT);
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS companies (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    code TEXT UNIQUE NOT NULL, name TEXT NOT NULL,
    address TEXT, address2 TEXT, address3 TEXT, address4 TEXT,
    telephone TEXT, email TEXT, trn TEXT, website TEXT,
    logo_path TEXT, stamp_path TEXT, letterhead_path TEXT,
    is_active INTEGER DEFAULT 1, sort_order INTEGER DEFAULT 0,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS clients (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL, company_name TEXT, trn TEXT,
    address TEXT, country TEXT, telephone TEXT, email TEXT, notes TEXT,
    is_active INTEGER DEFAULT 1,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS signatories (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL, position TEXT, image_path TEXT,
    is_active INTEGER DEFAULT 1, is_default INTEGER DEFAULT 0,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS stamps (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL, company_id INTEGER, image_path TEXT,
    is_active INTEGER DEFAULT 1, is_default INTEGER DEFAULT 0,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS drivers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL, contact TEXT, vehicle_number TEXT,
    is_active INTEGER DEFAULT 1,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS delivery_notes (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    dn_number TEXT UNIQUE NOT NULL,
    company_id INTEGER,
    dn_date TEXT NOT NULL,
    invoice_no TEXT, order_no TEXT,
    project_name TEXT, location TEXT,
    client_id INTEGER, client_name TEXT, client_telephone TEXT,
    invoice_address TEXT, shipping_address TEXT,
    show_prices INTEGER DEFAULT 0,
    currency TEXT DEFAULT 'AED',
    vat_rate REAL DEFAULT 5,
    subtotal REAL DEFAULT 0, vat_amount REAL DEFAULT 0, net_total REAL DEFAULT 0,
    vehicle_number TEXT, driver_name TEXT, driver_contact TEXT,
    notes TEXT,
    signatory_id INTEGER, signatory_name TEXT, signatory_position TEXT,
    signatory_phone TEXT, signatory_image TEXT,
    stamp_id INTEGER, stamp_name TEXT, stamp_image TEXT,
    include_signature INTEGER DEFAULT 1, include_stamp INTEGER DEFAULT 1,
    delivery_status TEXT DEFAULT 'pending',
    delivered_on TEXT, received_by TEXT, status_remarks TEXT,
    pdf_path TEXT, status TEXT DEFAULT 'active',
    created_by INTEGER,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS dn_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    dn_id INTEGER NOT NULL, sr_no INTEGER,
    description TEXT, quantity REAL DEFAULT 1, unit TEXT DEFAULT 'PCS',
    amount REAL DEFAULT 0, total_amount REAL DEFAULT 0);
CREATE TABLE IF NOT EXISTS audit_logs (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id INTEGER, username TEXT,
    action TEXT, details TEXT, ip_address TEXT,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP)
"""

def init_db():
    conn = get_db()
    conn.executescript(SCHEMA)
    conn.commit()

    # Safe column migrations (for databases created by an earlier version)
    for _t, _col, _typ in [("companies", "dn_prefix", "TEXT"),
                           ("signatories", "phone", "TEXT"),
                           ("users", "permissions", "TEXT"),
                           ("users", "driver_id", "INTEGER"),
                           ("delivery_notes", "map_link", "TEXT"),
                           ("delivery_notes", "driver_id", "INTEGER"),
                           ("delivery_notes", "delivered_by", "TEXT"),
                           ("delivery_notes", "delivered_at", "TEXT"),
                           ("delivery_notes", "task_type", "TEXT DEFAULT 'delivery'"),
                           ("delivery_notes", "sort_order", "INTEGER DEFAULT 0"),
                           ("delivery_notes", "pickup_name", "TEXT"),
                           ("delivery_notes", "pickup_phone", "TEXT"),
                           ("delivery_notes", "pickup_address", "TEXT"),
                           ("delivery_notes", "pickup_map_link", "TEXT"),
                           ("delivery_notes", "picked_up_at", "TEXT"),
                           ("delivery_notes", "picked_up_by", "TEXT"),
                           ("delivery_notes", "pickup_remarks", "TEXT"),
                           ("clients", "shipping_address", "TEXT"),
                           ("clients", "map_link", "TEXT")]:
        try:
            if USE_PG:
                conn.execute(f"ALTER TABLE {_t} ADD COLUMN IF NOT EXISTS {_col} {_typ}", ())
            else:
                conn.execute(f"ALTER TABLE {_t} ADD COLUMN {_col} {_typ}", ())
            conn.commit()
        except Exception:
            try: conn._conn.rollback()
            except Exception: pass

    # existing notes get a starting priority (oldest first); new notes go to the end of the line
    conn.execute("UPDATE delivery_notes SET sort_order=id WHERE sort_order IS NULL OR sort_order=0", ())
    conn.execute("UPDATE delivery_notes SET task_type='delivery' WHERE task_type IS NULL OR task_type=''", ())
    conn.commit()

    defaults = {
        'currency': 'AED',
        'logo_path': 'uploads/logos/GH_LOGO.jpg',
        'footer_path': 'images/gh_footer_AE.png',
        'default_vat_rate': '5',
        'default_company_id': '1',
        'default_dn_notes': DEFAULT_DN_NOTES,
    }
    for k, v in defaults.items():
        if USE_PG:
            conn.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO NOTHING", (k, v))
        else:
            conn.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k, v))

    if not conn.execute("SELECT id FROM users WHERE username=?", ('admin',)).fetchone():
        conn.execute("INSERT INTO users(username,password_hash,full_name,role) VALUES(?,?,?,?)",
                     ('admin', generate_password_hash(os.environ.get('ADMIN_PASSWORD') or 'admin123'), 'Administrator', 'admin'))

    companies_data = [
        ('GHM', 'Geometry Home Furniture Manufacturing L.L.C',
         'L06, L07, L08, WH Phase 3, Dubai Industrial City',
         'PO. Box: 16794,  Phone: +97145578898', '', '',
         '+97145578898', 'info@geometry-home.ae', '104150505600003', 'www.geometry-home.ae',
         'uploads/logos/GH_LOGO.jpg', 'uploads/stamps/GHM_Seal_transparent.png', '', 'GHM-DN', 1, 1),
        ('GHT', 'Geometry Home for Furniture Trading Co. L.L.C',
         'Shop - GF - 11, Al Barsha Second, Art of Living Mall, Dubai',
         'PO. Box: 16794,  Phone: +97148834020', '', '',
         '+97148834020', 'info@geometry-home.ae', '104150505600003', 'www.geometry-home.ae',
         'uploads/logos/GH_LOGO.jpg', 'uploads/stamps/GHT_Seal_transparent.png', '', 'GHT-DN', 1, 2),
        ('TRA', 'T R A Geometry Home Technical Services L.L.C S.O.C',
         'A25 \u2013 Hall No. 1, Al Nasr Central, Oud Metha, Dubai',
         'PO. Box: 16794,  Phone: +97145578898', '', '',
         '+97145578898', 'info@geometry-home.ae', '1533428', 'www.geometry-home.ae',
         'uploads/logos/GH_LOGO.jpg', 'uploads/stamps/TRA_Seal_transparent.png', '', 'TRA-DN', 1, 3),
    ]
    for cd in companies_data:
        if not conn.execute("SELECT id FROM companies WHERE code=?", (cd[0],)).fetchone():
            conn.execute("""INSERT INTO companies(code,name,address,address2,address3,address4,
                telephone,email,trn,website,logo_path,stamp_path,letterhead_path,
                dn_prefix,is_active,sort_order) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", cd)
    # companies created before the prefix existed
    for code in ('GHM', 'GHT', 'TRA'):
        conn.execute("UPDATE companies SET dn_prefix=? WHERE code=? AND (dn_prefix IS NULL OR dn_prefix='')",
                     (f'{code}-DN', code))

    if not conn.execute("SELECT id FROM signatories WHERE name=?", ('Rashid Khakimov',)).fetchone():
        conn.execute("INSERT INTO signatories(name,position,phone,image_path,is_active,is_default) VALUES(?,?,?,?,1,1)",
                     ('Rashid Khakimov', 'General Manager', '971 50 394 9037',
                      'uploads/signatures/Rashid_Khakimov_Sign_transparent.png'))

    for sname, scode, spath, isdef in [
            ('GHM Official Seal', 'GHM', 'uploads/stamps/GHM_Seal_transparent.png', 1),
            ('GHT Official Seal', 'GHT', 'uploads/stamps/GHT_Seal_transparent.png', 0),
            ('TRA Seal', 'TRA', 'uploads/stamps/TRA_Seal_transparent.png', 0)]:
        co_row = conn.execute("SELECT id FROM companies WHERE code=?", (scode,)).fetchone()
        s_row = conn.execute("SELECT id FROM stamps WHERE name=?", (sname,)).fetchone()
        if co_row and not s_row:
            conn.execute("INSERT INTO stamps(name,company_id,image_path,is_active,is_default) VALUES(?,?,?,1,?)",
                         (sname, co_row['id'], spath, isdef))
    conn.commit()
    conn.close()


# ── HELPERS ────────────────────────────────────────────────────────────────
def gset(key, default=''):
    conn = get_db()
    r = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    conn.close()
    return r['value'] if r else default

def gall():
    conn = get_db()
    rows = conn.execute("SELECT key,value FROM settings", ()).fetchall()
    conn.close()
    return {r['key']: r['value'] for r in rows}

def sset(key, value):
    conn = get_db()
    if USE_PG:
        conn.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=EXCLUDED.value", (key, value))
    else:
        conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (key, value))
    conn.commit(); conn.close()

def get_companies():
    conn = get_db()
    rows = conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order,code", ()).fetchall()
    conn.close()
    return rows

def get_company(cid):
    conn = get_db()
    row = conn.execute("SELECT * FROM companies WHERE id=?", (cid,)).fetchone()
    conn.close()
    return dict(row) if row else None

def audit(action, details=''):
    if 'user_id' in session:
        conn = get_db()
        conn.execute("INSERT INTO audit_logs(user_id,username,action,details,ip_address) VALUES(?,?,?,?,?)",
                     (session['user_id'], session.get('username'), action, details, request.remote_addr))
        conn.commit(); conn.close()

def n2w(amount):
    if amount is None: return ''
    try: amount = float(amount)
    except Exception: return ''
    if amount == 0: return 'Zero Dirhams Only'
    ones = ['', 'One', 'Two', 'Three', 'Four', 'Five', 'Six', 'Seven', 'Eight', 'Nine', 'Ten',
            'Eleven', 'Twelve', 'Thirteen', 'Fourteen', 'Fifteen', 'Sixteen', 'Seventeen', 'Eighteen', 'Nineteen']
    tens_w = ['', '', 'Twenty', 'Thirty', 'Forty', 'Fifty', 'Sixty', 'Seventy', 'Eighty', 'Ninety']
    def sp(n):
        n = int(n)
        if n < 20: return ones[n]
        if n < 100: return tens_w[n // 10] + (' ' + ones[n % 10] if n % 10 else '')
        if n < 1000: return ones[n // 100] + ' Hundred' + (' and ' + sp(n % 100) if n % 100 else '')
        if n < 1000000: return sp(n // 1000) + ' Thousand' + (' ' + sp(n % 1000) if n % 1000 else '')
        return sp(n // 1000000) + ' Million' + (' ' + sp(n % 1000000) if n % 1000000 else '')
    aed = int(amount); fils = round((amount - aed) * 100)
    r = sp(aed) if aed else 'Zero'
    if fils: r += ' and ' + sp(fils) + ' Fils'
    return r + ' Dirhams Only'

def fmt_date(d):
    if not d: return ''
    try:
        if '-' in str(d) and len(str(d)) == 10:
            return datetime.strptime(str(d), '%Y-%m-%d').strftime('%d-%m-%Y')
    except Exception: pass
    return str(d)

def fmt_date_long(d):
    try: return datetime.strptime(str(d), '%Y-%m-%d').strftime('%d-%b-%Y')
    except Exception: return str(d or '')

def fmt_qty(q):
    try: q = float(q)
    except Exception: return str(q or '')
    return str(int(q)) if q == int(q) else f"{q:g}"

def num(v, default=0.0):
    try: return float(str(v).replace(',', '').strip() or default)
    except Exception: return default

app.jinja_env.globals['fmt_date'] = fmt_date
app.jinja_env.globals['fmt_qty'] = fmt_qty

def nav_url(n):
    """Link a driver taps to navigate: the saved Google Map link, else a Maps search of the shipping address."""
    from urllib.parse import quote
    if n['map_link']: return n['map_link']
    addr = ' '.join((n['shipping_address'] or n['invoice_address'] or '').split())
    return f"https://www.google.com/maps/search/?api=1&query={quote(addr)}" if addr else ''
app.jinja_env.globals['nav_url'] = nav_url

def pickup_nav_url(n):
    """Navigate link for the collection point of a pickup job."""
    from urllib.parse import quote
    if n['pickup_map_link']: return n['pickup_map_link']
    addr = ' '.join((n['pickup_address'] or '').split())
    return f"https://www.google.com/maps/search/?api=1&query={quote(addr)}" if addr else ''
app.jinja_env.globals['pickup_nav_url'] = pickup_nav_url

def fmt_dt(value):
    """'2026-10-10 14:32' (already Dubai time) -> '10-Oct-2026 14:32'. Date-only values give just the date."""
    if not value: return ''
    t = str(value).strip()
    for f, out in (('%Y-%m-%d %H:%M:%S', '%d-%b-%Y %H:%M'), ('%Y-%m-%d %H:%M', '%d-%b-%Y %H:%M'), ('%Y-%m-%d', '%d-%b-%Y')):
        try: return datetime.strptime(t, f).strftime(out)
        except Exception: pass
    return t
app.jinja_env.filters['fmt_dt'] = fmt_dt

def is_pickup(n):
    try: return (n['task_type'] or 'delivery') == 'pickup'
    except Exception: return False
app.jinja_env.globals['is_pickup'] = is_pickup

DUBAI_TZ = timezone(timedelta(hours=4))

def now_dubai():
    return datetime.now(DUBAI_TZ).replace(tzinfo=None)

def dubai_time(value, fmt='%Y-%m-%d %H:%M'):
    """Show a stored UTC database timestamp in Dubai time (UTC+4)."""
    if not value: return '—'
    txt = str(value).strip()
    m = re.match(r'(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}:\d{2})(?:\.\d+)?\s*(Z|[+-]\d{2}(?::?\d{2})?)?$', txt)
    if not m: return txt[:19]
    try:
        dt = datetime.strptime(f"{m.group(1)} {m.group(2)}", '%Y-%m-%d %H:%M:%S')
        off = m.group(3)
        if off and off != 'Z':
            sign = -1 if off[0] == '-' else 1
            digits = off[1:].replace(':', '')
            offset = timedelta(hours=int(digits[:2]), minutes=int(digits[2:4] or 0)) * sign
        else:
            offset = timedelta(0)
        return (dt - offset + timedelta(hours=4)).strftime(fmt)
    except Exception:
        return txt[:19]

app.jinja_env.filters['dubai'] = dubai_time

def next_dn_number(company_id, task_type='delivery'):
    """Next free number for the company: PREFIX-YEAR-0001, 0002 ... (deleted notes keep their number).
    Pickup jobs use their own series: GHM-PU-2026-0001."""
    co = get_company(company_id) or {}
    prefix = (co.get('dn_prefix') or f"{co.get('code', 'GH')}-DN").strip()
    if task_type == 'pickup':
        prefix = (prefix[:-3] if prefix.upper().endswith('-DN') else prefix) + '-PU'
    year = now_dubai().year
    stem = f"{prefix}-{year}-"
    conn = get_db()
    rows = conn.execute("SELECT dn_number FROM delivery_notes WHERE dn_number LIKE ?", (stem + '%',)).fetchall()
    conn.close()
    top = 0
    for r in rows:
        try: top = max(top, int(r['dn_number'].rsplit('-', 1)[-1]))
        except Exception: pass
    return f"{stem}{top + 1:04d}"


# ── PERMISSIONS ────────────────────────────────────────────────────────────
# Admin: everything.  Delivery user (role 'user'): delivery notes + drivers, with per-user switches below.
DN_PERMS = {
    'dn_create':   'Create delivery notes (new and duplicate)',
    'dn_edit':     'Edit delivery notes',
    'dn_deliver':  'Mark delivery notes as delivered (and update delivery status)',
    'dn_delete':   'Delete delivery notes',
    'dn_view_all': "See and manage other users' delivery notes (otherwise only their own)",
}
DEFAULT_PERMS = ['dn_create', 'dn_edit', 'dn_deliver', 'dn_view_all']
USER_ENDPOINTS = {
    'index', 'login', 'logout', 'change_password', 'static',
    'notes', 'new_note', 'edit_note', 'view_note', 'delete_note', 'duplicate_note',
    'download_pdf', 'preview_pdf', 'set_status', 'reorder_notes',
    'drivers', 'save_driver', 'delete_driver',
    'api_clients', 'api_client_save', 'clients', 'new_client', 'edit_client', 'api_stamps', 'api_next_number', 'api_drivers',
}
ENDPOINT_PERM = {
    'new_note': 'dn_create', 'duplicate_note': 'dn_create',
    'edit_note': 'dn_edit', 'set_status': 'dn_deliver',
    'delete_note': 'dn_delete', 'reorder_notes': 'dn_edit',
}

def user_perms(role, raw):
    if role == 'admin': return set(DN_PERMS)
    if raw is None: return set(DEFAULT_PERMS)
    return {p for p in str(raw).split(',') if p in DN_PERMS}

DRIVER_ENDPOINTS = {'driver_home', 'driver_deliver', 'driver_pickup', 'driver_password', 'driver_logout', 'logout', 'index', 'uploaded_file'}

@app.before_request
def enforce_permissions():
    """Secure by default: anything not on the user allow-list is admin-only."""
    ep = request.endpoint
    if ep in (None, 'static', 'login', 'driver_login', 'health', 'uploaded_file') or 'user_id' not in session: return
    conn = get_db()
    u = conn.execute("SELECT id,role,is_active,permissions,driver_id FROM users WHERE id=?", (session['user_id'],)).fetchone()
    conn.close()
    if not u or not u['is_active']:
        session.clear(); flash('Your account is not active.', 'danger'); return redirect(url_for('login'))
    session['role'] = u['role']
    g.perms = user_perms(u['role'], u['permissions'])
    g.driver_id = u['driver_id']
    if u['role'] == 'driver':                      # drivers only ever see their own portal
        if ep not in DRIVER_ENDPOINTS: return redirect(url_for('driver_home'))
        return
    if u['role'] == 'admin': return
    need = ENDPOINT_PERM.get(ep)
    if ep not in USER_ENDPOINTS or (need and need not in g.perms):
        flash('You do not have permission to do that.', 'danger')
        return redirect(url_for('notes'))

@app.context_processor
def inject_perms():
    def can(p): return session.get('role') == 'admin' or p in getattr(g, 'perms', ())
    return dict(can=can)

def can_see_all():
    return session.get('role') == 'admin' or 'dn_view_all' in getattr(g, 'perms', ())

def note_allowed(nid):
    if can_see_all(): return True
    conn = get_db(); r = conn.execute("SELECT created_by FROM delivery_notes WHERE id=?", (nid,)).fetchone(); conn.close()
    return bool(r) and r['created_by'] == session.get('user_id')

def _no_access():
    flash('You do not have access to that delivery note.', 'danger')
    return redirect(url_for('notes'))

def home_url():
    if session.get('role') == 'admin': return url_for('dashboard')
    if session.get('role') == 'driver': return url_for('driver_home')
    return url_for('notes')

def login_req(f):
    @wraps(f)
    def d(*a, **k):
        if 'user_id' not in session: return redirect(url_for('login'))
        return f(*a, **k)
    return d

def driver_req(f):
    @wraps(f)
    def d(*a, **k):
        if 'user_id' not in session: return redirect(url_for('driver_login'))
        if session.get('role') != 'driver':
            flash('The driver portal is for driver accounts only.', 'danger')
            return redirect(home_url())
        return f(*a, **k)
    return d

def admin_req(f):
    @wraps(f)
    def d(*a, **k):
        if 'user_id' not in session: return redirect(url_for('login'))
        if session.get('role') != 'admin':
            flash('Admin access required.', 'danger')
            return redirect(url_for('notes'))
        return f(*a, **k)
    return d


# ── AUTH ───────────────────────────────────────────────────────────────────
@app.route('/')
def index(): return redirect(home_url() if 'user_id' in session else url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        u = request.form.get('username', '').strip(); p = request.form.get('password', '')
        conn = get_db(); user = conn.execute("SELECT * FROM users WHERE username=? AND is_active=1", (u,)).fetchone(); conn.close()
        if user and check_password_hash(user['password_hash'], p):
            session.clear()
            session['user_id'] = user['id']; session['username'] = user['username']
            session['full_name'] = user['full_name'] or user['username']; session['role'] = user['role']
            session.permanent = 'remember' in request.form
            conn2 = get_db(); conn2.execute("UPDATE users SET last_login=CURRENT_TIMESTAMP WHERE id=?", (user['id'],)); conn2.commit(); conn2.close()
            audit('LOGIN'); return redirect(home_url())
        flash('Invalid username or password.', 'danger')
    return render_template('login.html', s=gall())

@app.route('/logout')
def logout():
    was_driver = session.get('role') == 'driver'
    audit('LOGOUT'); session.clear()
    return redirect(url_for('driver_login' if was_driver else 'login'))

@app.route('/change-password', methods=['GET', 'POST'])
@login_req
def change_password():
    if request.method == 'POST':
        cur = request.form.get('current_password', ''); nw = request.form.get('new_password', ''); cn = request.form.get('confirm_password', '')
        if nw != cn: flash('Passwords do not match.', 'danger')
        elif len(nw) < 6: flash('Minimum 6 characters.', 'danger')
        else:
            conn = get_db(); user = conn.execute("SELECT * FROM users WHERE id=?", (session['user_id'],)).fetchone()
            if user and check_password_hash(user['password_hash'], cur):
                conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(nw), session['user_id'])); conn.commit()
                flash('Password changed.', 'success')
            else: flash('Current password incorrect.', 'danger')
            conn.close()
    return render_template('change_password.html', s=gall())


# ── DASHBOARD (admin) ──────────────────────────────────────────────────────
@app.route('/dashboard')
@admin_req
def dashboard():
    conn = get_db()
    def count(extra='', params=()):
        return conn.execute(f"SELECT COUNT(*) c FROM delivery_notes WHERE status='active'{extra}", params).fetchone()['c']
    mo = now_dubai().strftime('%Y-%m')
    stats = {
        'total': count(),
        'pending': count(" AND delivery_status='pending'"),
        'delivered': count(" AND delivery_status='delivered'"),
        'month': count(" AND dn_date LIKE ?", (f'{mo}%',)),
        'month_delivered': count(" AND delivery_status='delivered' AND dn_date LIKE ?", (f'{mo}%',)),
    }
    recent = conn.execute("""SELECT n.*,c.code as co_code FROM delivery_notes n
        LEFT JOIN companies c ON n.company_id=c.id WHERE n.status='active' ORDER BY n.id DESC LIMIT 10""").fetchall()
    old_pending = conn.execute("""SELECT n.*,c.code as co_code FROM delivery_notes n
        LEFT JOIN companies c ON n.company_id=c.id
        WHERE n.status='active' AND n.delivery_status='pending' AND n.dn_date<=?
        ORDER BY n.dn_date LIMIT 10""", ((now_dubai() - timedelta(days=3)).strftime('%Y-%m-%d'),)).fetchall()
    companies = conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    co_stats = {}
    for co in companies:
        co_stats[co['id']] = {
            'total': count(" AND company_id=?", (co['id'],)),
            'pending': count(" AND delivery_status='pending' AND company_id=?", (co['id'],))}
    # ── extra statistics ──
    today = now_dubai().strftime('%Y-%m-%d')
    overdue_cut = (now_dubai() - timedelta(days=3)).strftime('%Y-%m-%d')
    dd = "COALESCE(NULLIF(delivered_on,''), SUBSTR(COALESCE(delivered_at,''),1,10))"   # date the job was actually delivered
    stats['pending_month'] = count(" AND delivery_status='pending' AND dn_date LIKE ?", (f'{mo}%',))
    stats['delivered_month'] = conn.execute(
        f"SELECT COUNT(*) c FROM delivery_notes WHERE status='active' AND delivery_status='delivered' AND {dd} LIKE ?",
        (f'{mo}%',)).fetchone()['c']
    stats['delivered_today'] = conn.execute(
        f"SELECT COUNT(*) c FROM delivery_notes WHERE status='active' AND delivery_status='delivered' AND {dd}=?",
        (today,)).fetchone()['c']
    stats['pending_today'] = count(" AND delivery_status='pending' AND dn_date=?", (today,))
    stats['overdue'] = count(" AND delivery_status='pending' AND dn_date<=?", (overdue_cut,))
    stats['cancelled_month'] = count(" AND delivery_status='cancelled' AND dn_date LIKE ?", (f'{mo}%',))
    month_jobs = stats['pending_month'] + stats['month_delivered']
    stats['month_rate'] = round(100 * stats['month_delivered'] / month_jobs) if month_jobs else 0

    # per driver (grouped by name; blank = Unassigned)
    drv_sql = f"""SELECT COALESCE(NULLIF(TRIM(driver_name),''),'Unassigned') AS drv,
        SUM(CASE WHEN delivery_status='pending' THEN 1 ELSE 0 END) AS pending,
        SUM(CASE WHEN delivery_status='pending' AND dn_date LIKE ? THEN 1 ELSE 0 END) AS pending_month,
        SUM(CASE WHEN delivery_status='pending' AND dn_date<=? THEN 1 ELSE 0 END) AS overdue,
        SUM(CASE WHEN delivery_status='delivered' AND {dd}=? THEN 1 ELSE 0 END) AS delivered_today,
        SUM(CASE WHEN delivery_status='delivered' AND {dd} LIKE ? THEN 1 ELSE 0 END) AS delivered_month,
        SUM(CASE WHEN delivery_status='delivered' THEN 1 ELSE 0 END) AS delivered_total
        FROM delivery_notes WHERE status='active' AND delivery_status IN ('pending','delivered')
        GROUP BY COALESCE(NULLIF(TRIM(driver_name),''),'Unassigned')"""
    driver_stats = [dict(r) for r in conn.execute(drv_sql, (f'{mo}%', overdue_cut, today, f'{mo}%')).fetchall()]
    for r in driver_stats:
        for k in ('pending', 'pending_month', 'overdue', 'delivered_today', 'delivered_month', 'delivered_total'):
            r[k] = int(r[k] or 0)
    driver_stats.sort(key=lambda r: (-r['pending'], -r['delivered_month'], r['drv'].lower()))
    driver_tot = {k: sum(r[k] for r in driver_stats) for k in
                  ('pending', 'pending_month', 'overdue', 'delivered_today', 'delivered_month', 'delivered_total')}

    # last 6 months: jobs created vs delivered vs still pending
    mrows = conn.execute("""SELECT SUBSTR(dn_date,1,7) AS m, COUNT(*) AS total,
        SUM(CASE WHEN delivery_status='delivered' THEN 1 ELSE 0 END) AS delivered,
        SUM(CASE WHEN delivery_status='pending' THEN 1 ELSE 0 END) AS pending
        FROM delivery_notes WHERE status='active' AND dn_date IS NOT NULL AND dn_date<>''
        GROUP BY SUBSTR(dn_date,1,7) ORDER BY m DESC LIMIT 6""").fetchall()
    monthly = [dict(m=r['m'], total=int(r['total'] or 0), delivered=int(r['delivered'] or 0), pending=int(r['pending'] or 0)) for r in mrows]
    logs = conn.execute("SELECT * FROM audit_logs ORDER BY id DESC LIMIT 8").fetchall()
    conn.close()
    return render_template('dashboard.html', s=gall(), stats=stats, recent=recent, old_pending=old_pending,
                           companies=companies, co_stats=co_stats, logs=logs, driver_stats=driver_stats,
                           driver_tot=driver_tot, monthly=monthly, month_label=now_dubai().strftime('%B %Y'))


# ── DELIVERY NOTES ─────────────────────────────────────────────────────────
@app.route('/notes')
@login_req
def notes():
    view = 'list' if request.args.get('view') == 'list' else 'driver'
    q = request.args.get('q', '').strip(); co_id = request.args.get('co', '')
    st = request.args.get('st', 'pending' if (view == 'driver' and 'st' not in request.args) else '')
    tp = request.args.get('tp', '')
    d_from = request.args.get('from', ''); d_to = request.args.get('to', '')
    sql = """SELECT n.*,c.code as co_code,
             (SELECT COUNT(*) FROM dn_items i WHERE i.dn_id=n.id) AS item_count
             FROM delivery_notes n LEFT JOIN companies c ON n.company_id=c.id WHERE n.status='active'"""
    params = []
    if q:
        like = f'%{q.lower()}%'
        sql += """ AND (LOWER(n.dn_number) LIKE ? OR LOWER(COALESCE(n.client_name,'')) LIKE ?
                   OR LOWER(COALESCE(n.project_name,'')) LIKE ? OR LOWER(COALESCE(n.invoice_no,'')) LIKE ?
                   OR LOWER(COALESCE(n.order_no,'')) LIKE ? OR LOWER(COALESCE(n.driver_name,'')) LIKE ?
                   OR LOWER(COALESCE(n.pickup_name,'')) LIKE ?)"""
        params += [like] * 7
    if co_id: sql += " AND n.company_id=?"; params.append(co_id)
    if st in ('pending', 'delivered', 'cancelled'): sql += " AND n.delivery_status=?"; params.append(st)
    if tp in ('delivery', 'pickup'): sql += " AND COALESCE(n.task_type,'delivery')=?"; params.append(tp)
    if d_from: sql += " AND n.dn_date>=?"; params.append(d_from)
    if d_to: sql += " AND n.dn_date<=?"; params.append(d_to)
    if not can_see_all(): sql += " AND n.created_by=?"; params.append(session.get('user_id'))
    if view == 'list': sql += " ORDER BY n.id DESC"
    else: sql += " ORDER BY n.dn_date, COALESCE(n.sort_order,0), n.id"
    conn = get_db()
    rows = conn.execute(sql, params).fetchall()
    companies = conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    drv_rows = conn.execute("SELECT * FROM drivers WHERE is_active=1 ORDER BY name").fetchall() if view == 'driver' else []
    conn.close()

    groups = []
    if view == 'driver':
        dates = sorted({r['dn_date'] for r in rows})
        by_drv = {}
        for r in rows: by_drv.setdefault(r['driver_id'], []).append(r)
        def build(did, name, vehicle, contact):
            cells = []
            for d in dates:
                mine = [r for r in by_drv.get(did, []) if r['dn_date'] == d]
                pend = [r for r in mine if r['delivery_status'] == 'pending']
                done = [r for r in mine if r['delivery_status'] != 'pending']
                cells.append({'date': d, 'pending': pend, 'done': done})
            allr = by_drv.get(did, [])
            return {'driver_id': did, 'name': name, 'vehicle': vehicle, 'contact': contact, 'cells': cells,
                    'pending': sum(1 for r in allr if r['delivery_status'] == 'pending'),
                    'pickups': sum(1 for r in allr if r['delivery_status'] == 'pending' and is_pickup(r)),
                    'total': len(allr)}
        groups.append(build(None, 'Unassigned', '', ''))
        known = set()
        for d in drv_rows:
            known.add(d['id']); groups.append(build(d['id'], d['name'], d['vehicle_number'] or '', d['contact'] or ''))
        for did in by_drv:                     # notes of drivers that were removed since
            if did is not None and did not in known:
                r0 = by_drv[did][0]
                groups.append(build(did, (r0['driver_name'] or 'Removed driver') + ' (removed)', r0['vehicle_number'] or '', ''))
        # keep quiet drivers visible only when they matter: hide empty groups when searching
        if q: groups = [g_ for g_ in groups if g_['total']]
        elif not dates: groups = [g_ for g_ in groups if g_['total']]
    return render_template('notes.html', notes=rows, groups=groups, view=view, q=q, co_id=co_id, st=st, tp=tp,
                           d_from=d_from, d_to=d_to, companies=companies, s=gall(), today=now_dubai().strftime('%Y-%m-%d'),
                           now_time=now_dubai().strftime('%H:%M'),
                           tomorrow=(now_dubai() + timedelta(days=1)).strftime('%Y-%m-%d'))


@app.route('/notes/reorder', methods=['POST'])
@login_req
def reorder_notes():
    """Priority board: drag a delivery/pickup up or down a driver's list, or onto another driver.
    Body: {"lists": [{"driver_id": 3 | null, "ids": [12, 9, 15]}, ...]} - each list is the new top-to-bottom order."""
    data = request.get_json(silent=True) or {}
    lists = data.get('lists') or []
    conn = get_db(); changed = 0; stale = []; moved_to = set()
    try:
        for lst in lists:
            did = _int_or_none(lst.get('driver_id'))
            drv = conn.execute("SELECT * FROM drivers WHERE id=? AND is_active=1", (did,)).fetchone() if did else None
            if did and not drv: continue
            for pos, raw in enumerate(lst.get('ids') or [], 1):
                nid = _int_or_none(raw)
                if not nid or not note_allowed(nid): continue
                n = conn.execute("SELECT id,driver_id,delivery_status,pdf_path,dn_number FROM delivery_notes WHERE id=? AND status='active'", (nid,)).fetchone()
                if not n or n['delivery_status'] != 'pending': continue
                if (n['driver_id'] or None) != did:
                    conn.execute("""UPDATE delivery_notes SET driver_id=?,driver_name=?,driver_contact=?,vehicle_number=?,sort_order=?,
                                    pdf_path=NULL,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                                 (did, drv['name'] if drv else '', (drv['contact'] if drv else '') or '',
                                  (drv['vehicle_number'] if drv else '') or '', pos, nid))
                    if n['pdf_path']: stale.append(n['pdf_path'])
                    moved_to.add(f"{n['dn_number']} → {drv['name'] if drv else 'Unassigned'}")
                else:
                    conn.execute("UPDATE delivery_notes SET sort_order=? WHERE id=?", (pos, nid))
                changed += 1
        conn.commit()
    finally:
        conn.close()
    for sp in stale:
        try:
            if os.path.exists(sp): os.remove(sp)
        except Exception: pass
    audit('DN_PRIORITY', '; '.join(sorted(moved_to)) if moved_to else f'{changed} reordered')
    return jsonify({'ok': True, 'changed': changed})


def _form_context(conn, co_id):
    companies = conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    clients = conn.execute("SELECT * FROM clients WHERE is_active=1 ORDER BY name").fetchall()
    sigs = conn.execute("SELECT * FROM signatories WHERE is_active=1 ORDER BY is_default DESC,name").fetchall()
    stamps = conn.execute("SELECT * FROM stamps WHERE is_active=1 AND company_id=? ORDER BY is_default DESC,name", (co_id,)).fetchall()
    if not stamps:
        stamps = conn.execute("SELECT * FROM stamps WHERE is_active=1 ORDER BY is_default DESC").fetchall()
    drv = conn.execute("SELECT * FROM drivers WHERE is_active=1 ORDER BY name").fetchall()
    return companies, clients, sigs, stamps, drv

@app.route('/notes/new', methods=['GET', 'POST'])
@login_req
def new_note():
    if request.method == 'POST': return _save_note(None)
    s = gall(); conn = get_db()
    try: co_id = int(request.args.get('co', s.get('default_company_id', '1')))
    except Exception: co_id = 1
    ttype = 'pickup' if request.args.get('type') == 'pickup' else 'delivery'
    companies, clients, sigs, stamps, drv = _form_context(conn, co_id)
    def_sig = conn.execute("SELECT * FROM signatories WHERE is_default=1 AND is_active=1 LIMIT 1").fetchone()
    def_stamp = conn.execute("SELECT * FROM stamps WHERE is_default=1 AND is_active=1 AND company_id=? LIMIT 1", (co_id,)).fetchone()
    conn.close()
    return render_template('note_form.html', mode='new', n=None, items=[], companies=companies, clients=clients,
                           sigs=sigs, stamps=stamps, drivers=drv, def_sig=def_sig, def_stamp=def_stamp,
                           dn_num=next_dn_number(co_id, ttype), selected_co_id=co_id, units=UNITS, s=s, ttype=ttype,
                           today=now_dubai().strftime('%Y-%m-%d'))

@app.route('/notes/<int:nid>/edit', methods=['GET', 'POST'])
@login_req
def edit_note(nid):
    if not note_allowed(nid): return _no_access()
    if request.method == 'POST': return _save_note(nid)
    s = gall(); conn = get_db()
    n = conn.execute("SELECT * FROM delivery_notes WHERE id=? AND status='active'", (nid,)).fetchone()
    if not n:
        conn.close(); flash('Not found.', 'danger'); return redirect(url_for('notes'))
    items = conn.execute("SELECT * FROM dn_items WHERE dn_id=? ORDER BY sr_no", (nid,)).fetchall()
    companies, clients, sigs, stamps, drv = _form_context(conn, n['company_id'] or 1)
    conn.close()
    return render_template('note_form.html', mode='edit', n=n, items=items, companies=companies, clients=clients,
                           sigs=sigs, stamps=stamps, drivers=drv, def_sig=None, def_stamp=None,
                           dn_num=n['dn_number'], selected_co_id=n['company_id'] or 1, units=UNITS, s=s, ttype=('pickup' if is_pickup(n) else 'delivery'),
                           today=now_dubai().strftime('%Y-%m-%d'))

def clean_map_link(raw):
    """Accepts a Google Maps link (any https URL) or plain 'lat,lng'. Returns (link, warning)."""
    v = (raw or '').strip()
    if not v: return '', ''
    m = re.match(r'^(-?\d{1,3}(?:\.\d+)?)\s*,\s*(-?\d{1,3}(?:\.\d+)?)$', v)
    if m:
        return f"https://www.google.com/maps/dir/?api=1&destination={m.group(1)},{m.group(2)}", ''
    if re.match(r'^https?://[^\s]+$', v, re.I) and len(v) <= 2000:
        return v, ''
    return '', 'The Google Map location was not a valid link, so it was not saved.'

def _int_or_none(v):
    try: return int(str(v).strip())
    except (TypeError, ValueError): return None

def _autosave_client(conn, name, tel='', addr='', ship='', map_link='', client_id=None, last_used=True, overwrite=True):
    """Remember a customer and keep their details current.
    - New name -> a client record is created.
    - Known client -> any detail typed on the note that differs from the saved one replaces it (so edits stick);
      blank fields on the note never erase saved details.
    - The client is matched by id only when the name still matches, so typing a different customer name never edits
      the previously selected client."""
    name = (name or '').strip()
    if not name: return client_id
    row = None
    if client_id:
        row = conn.execute("SELECT * FROM clients WHERE id=? AND is_active=1", (client_id,)).fetchone()
        if row and (row['name'] or '').strip().lower() != name.lower(): row = None
    if not row: row = conn.execute("SELECT * FROM clients WHERE LOWER(name)=LOWER(?) AND is_active=1 ORDER BY id LIMIT 1", (name,)).fetchone()
    vals = {'telephone': (tel or '').strip(), 'address': (addr or '').strip(),
            'shipping_address': (ship or '').strip(), 'map_link': (map_link or '').strip()}
    if row:
        sets, params = [], []
        for col, v in vals.items():
            if not v: continue
            cur = (row[col] or '').strip() if row[col] else ''
            force = overwrite if col in ('telephone', 'address') else (last_used or overwrite)
            if v != cur and (force or not cur):
                sets.append(f'{col}=?'); params.append(v)
        if sets: conn.execute(f"UPDATE clients SET {','.join(sets)} WHERE id=?", params + [row['id']])
        return row['id']
    conn.execute(returning_id("INSERT INTO clients(name,telephone,address,shipping_address,map_link) VALUES(?,?,?,?,?)"),
                 (name, vals['telephone'], vals['address'], vals['shipping_address'], vals['map_link']))
    return last_insert_id(conn)

def _autosave_driver(conn, name, contact='', vehicle=''):
    """Typed-in driver (not picked from the list) is saved automatically; returns the driver id."""
    name = (name or '').strip()
    if not name: return None
    contact, vehicle = (contact or '').strip(), (vehicle or '').strip()
    ex = conn.execute("SELECT id,contact FROM drivers WHERE LOWER(name)=LOWER(?) AND COALESCE(vehicle_number,'')=? AND is_active=1", (name, vehicle)).fetchone()
    if ex:
        if contact and not (ex['contact'] or '').strip(): conn.execute("UPDATE drivers SET contact=? WHERE id=?", (contact, ex['id']))
        return ex['id']
    conn.execute(returning_id("INSERT INTO drivers(name,contact,vehicle_number) VALUES(?,?,?)"), (name, contact, vehicle))
    return last_insert_id(conn)

def _save_note(nid):
    try:
        return _save_note_impl(nid)
    except Exception as e:
        app.logger.exception('Saving delivery note failed')
        flash(f'Could not save the delivery note: {e}', 'danger')
        return redirect(url_for('edit_note', nid=nid) if nid else url_for('new_note'))

def _save_note_impl(nid):
    f = request.form
    was_edit = bool(nid)
    descs = f.getlist('item_desc[]'); qtys = f.getlist('item_qty[]')
    units = f.getlist('item_unit[]'); amts = f.getlist('item_amount[]')
    task_type = 'pickup' if f.get('task_type') == 'pickup' else 'delivery'
    show_prices = 1 if (f.get('show_prices') and task_type == 'delivery') else 0
    vat_rate = num(f.get('vat_rate'), 5)

    items = []
    for i, d in enumerate(descs):
        d = (d or '').strip()
        if not d: continue
        q = num(qtys[i] if i < len(qtys) else 1, 1)
        u = ((units[i] if i < len(units) else '') or 'PCS').strip().upper()
        a = num(amts[i] if i < len(amts) else 0) if show_prices else 0.0
        items.append((d, q, u, a, round(q * a, 2)))
    if not items:
        flash('Add at least one item to the delivery note.', 'danger')
        return redirect(url_for('edit_note', nid=nid) if nid else url_for('new_note'))

    subtotal = round(sum(it[4] for it in items), 2) if show_prices else 0.0
    vat_amt = round(subtotal * vat_rate / 100, 2) if show_prices else 0.0
    net = round(subtotal + vat_amt, 2)

    co_id = _int_or_none(f.get('company_id')) or 1
    map_link, map_warn = clean_map_link(f.get('map_link'))
    pu_map, pu_warn = clean_map_link(f.get('pickup_map_link'))
    conn = get_db()
    try:
        drv_row = None
        if _int_or_none(f.get('driver_id')):
            drv_row = conn.execute("SELECT * FROM drivers WHERE id=? AND is_active=1", (_int_or_none(f.get('driver_id')),)).fetchone()
        if not drv_row and f.get('driver_name', '').strip():       # typed-in driver: remember automatically
            new_drv = _autosave_driver(conn, f.get('driver_name'), f.get('driver_contact'), f.get('vehicle_number'))
            drv_row = conn.execute("SELECT * FROM drivers WHERE id=?", (new_drv,)).fetchone() if new_drv else None
        sig = conn.execute("SELECT * FROM signatories WHERE id=?", (_int_or_none(f.get('signatory_id')),)).fetchone() \
            if _int_or_none(f.get('signatory_id')) else None
        stp = conn.execute("SELECT * FROM stamps WHERE id=?", (_int_or_none(f.get('stamp_id')),)).fetchone() \
            if _int_or_none(f.get('stamp_id')) else None
        client_id = _autosave_client(conn, f.get('client_name'), f.get('client_telephone'), f.get('invoice_address'),
                                     f.get('shipping_address'), map_link, client_id=_int_or_none(f.get('client_id')))
        if task_type == 'pickup':                                   # the party we collect from is remembered too
            _autosave_client(conn, f.get('pickup_name'), f.get('pickup_phone'), f.get('pickup_address'), '', pu_map, last_used=False)
        notes_txt = (f.get('notes') or '').strip()
        base = [co_id, f.get('dn_date') or now_dubai().strftime('%Y-%m-%d'),
                f.get('invoice_no', '').strip(), f.get('order_no', '').strip(),
                f.get('project_name', '').strip(), f.get('location', '').strip(),
                client_id, f.get('client_name', '').strip(), f.get('client_telephone', '').strip(),
                f.get('invoice_address', '').strip(), f.get('shipping_address', '').strip(),
                show_prices, gset('currency', 'AED'), vat_rate, subtotal, vat_amt, net,
                f.get('vehicle_number', '').strip() or (drv_row['vehicle_number'] if drv_row else '') or '',
                f.get('driver_name', '').strip() or (drv_row['name'] if drv_row else ''),
                f.get('driver_contact', '').strip() or (drv_row['contact'] if drv_row else '') or '',
                notes_txt,
                sig['id'] if sig else None, sig['name'] if sig else '', sig['position'] if sig else '',
                (sig['phone'] if sig and 'phone' in sig.keys() else '') or '', sig['image_path'] if sig else '',
                stp['id'] if stp else None, stp['name'] if stp else '', stp['image_path'] if stp else '',
                1 if f.get('include_signature') else 0, 1 if f.get('include_stamp') else 0,
                map_link, drv_row['id'] if drv_row else None,
                task_type, f.get('pickup_name', '').strip() if task_type == 'pickup' else '',
                f.get('pickup_phone', '').strip() if task_type == 'pickup' else '',
                f.get('pickup_address', '').strip() if task_type == 'pickup' else '',
                pu_map if task_type == 'pickup' else '']
        cols = ("company_id,dn_date,invoice_no,order_no,project_name,location,client_id,client_name,client_telephone,"
                "invoice_address,shipping_address,show_prices,currency,vat_rate,subtotal,vat_amount,net_total,"
                "vehicle_number,driver_name,driver_contact,notes,signatory_id,signatory_name,signatory_position,"
                "signatory_phone,signatory_image,stamp_id,stamp_name,stamp_image,include_signature,include_stamp,"
                "map_link,driver_id,task_type,pickup_name,pickup_phone,pickup_address,pickup_map_link").split(',')

        if nid:
            old = conn.execute("SELECT dn_number,pdf_path FROM delivery_notes WHERE id=?", (nid,)).fetchone()
            if not old: raise ValueError('Delivery note not found')
            sets = ','.join(f"{c}=?" for c in cols) + ",pdf_path=NULL,updated_at=CURRENT_TIMESTAMP"
            conn.execute(f"UPDATE delivery_notes SET {sets} WHERE id=?", base + [nid])
            conn.execute("DELETE FROM dn_items WHERE dn_id=?", (nid,))
            number = old['dn_number']; stale = old['pdf_path']
        else:
            stale = None
            for _ in range(5):                      # retry if two people grab the same number at once
                number = next_dn_number(co_id, task_type)
                try:
                    top = conn.execute("SELECT COALESCE(MAX(sort_order),0) AS m FROM delivery_notes", ()).fetchone()['m']
                    sql = (f"INSERT INTO delivery_notes(dn_number,{','.join(cols)},sort_order,created_by) "
                           f"VALUES({','.join(['?'] * (len(cols) + 3))})")
                    conn.execute(returning_id(sql), [number] + base + [int(top) + 1, session['user_id']])
                    nid = last_insert_id(conn)
                    break
                except Exception as e:
                    try: conn._conn.rollback()
                    except Exception: pass
                    if 'unique' not in str(e).lower() and 'duplicate' not in str(e).lower(): raise
            else:
                raise RuntimeError('Could not allocate a delivery note number, please try again.')
        for i, it in enumerate(items, 1):
            conn.execute("INSERT INTO dn_items(dn_id,sr_no,description,quantity,unit,amount,total_amount) VALUES(?,?,?,?,?,?,?)",
                         (nid, i) + it)
        conn.commit()
    finally:
        conn.close()
    if stale and os.path.exists(stale):
        try: os.remove(stale)
        except Exception: pass
    audit('DN_EDITED' if was_edit else 'DN_CREATED', number)
    flash(f"{'Pickup' if task_type == 'pickup' else 'Delivery note'} {number} saved.", 'success')
    if map_warn: flash(map_warn, 'danger')
    if pu_warn and task_type == 'pickup': flash(pu_warn.replace('Google Map location', 'Pickup Google Map location'), 'danger')
    return redirect(url_for('view_note', nid=nid))

@app.route('/notes/<int:nid>')
@login_req
def view_note(nid):
    if not note_allowed(nid): return _no_access()
    conn = get_db()
    n = conn.execute("""SELECT n.*,c.name as co_name,c.code as co_code FROM delivery_notes n
        LEFT JOIN companies c ON n.company_id=c.id WHERE n.id=? AND n.status='active'""", (nid,)).fetchone()
    if not n:
        conn.close(); flash('Not found.', 'danger'); return redirect(url_for('notes'))
    items = conn.execute("SELECT * FROM dn_items WHERE dn_id=? ORDER BY sr_no", (nid,)).fetchall()
    conn.close()
    return render_template('note_view.html', n=n, items=items, s=gall(), today=now_dubai().strftime('%Y-%m-%d'),
                           now_time=now_dubai().strftime('%H:%M'))

@app.route('/notes/<int:nid>/status', methods=['POST'])
@login_req
def set_status(nid):
    if not note_allowed(nid): return _no_access()
    st = request.form.get('delivery_status', 'pending')
    if st not in ('pending', 'delivered', 'cancelled'): st = 'pending'
    delivered_on = request.form.get('delivered_on') or now_dubai().strftime('%Y-%m-%d')
    dtime = (request.form.get('delivered_time') or '').strip()
    if not re.match(r'^\d{2}:\d{2}$', dtime): dtime = now_dubai().strftime('%H:%M')
    conn = get_db()
    row = conn.execute("SELECT dn_number FROM delivery_notes WHERE id=?", (nid,)).fetchone()
    if row:
        if st == 'delivered':
            conn.execute("""UPDATE delivery_notes SET delivery_status=?,delivered_on=?,received_by=?,status_remarks=?,
                            delivered_by=?,delivered_at=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                         (st, delivered_on, request.form.get('received_by', '').strip(), request.form.get('status_remarks', '').strip(),
                          session.get('full_name') or session.get('username'), f"{delivered_on} {dtime}", nid))
        else:
            conn.execute("""UPDATE delivery_notes SET delivery_status=?,delivered_on=NULL,received_by=NULL,status_remarks=?,
                            delivered_by=NULL,delivered_at=NULL,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                         (st, request.form.get('status_remarks', '').strip(), nid))
        conn.commit()
        audit('DN_STATUS', f"{row['dn_number']} → {st}")
        flash(f"{row['dn_number']} marked as {st}.", 'success')
    conn.close()
    if request.form.get('next') == 'list': return redirect(url_for('notes'))
    return redirect(url_for('view_note', nid=nid))

@app.route('/notes/<int:nid>/delete', methods=['POST'])
@login_req
def delete_note(nid):
    if not note_allowed(nid): return _no_access()
    conn = get_db(); n = conn.execute("SELECT dn_number FROM delivery_notes WHERE id=?", (nid,)).fetchone()
    conn.execute("UPDATE delivery_notes SET status='deleted' WHERE id=?", (nid,)); conn.commit(); conn.close()
    if n: audit('DN_DELETED', n['dn_number'])
    flash('Delivery note deleted.', 'success'); return redirect(url_for('notes'))

@app.route('/notes/<int:nid>/duplicate')
@login_req
def duplicate_note(nid):
    if not note_allowed(nid): return _no_access()
    conn = get_db()
    n = conn.execute("SELECT * FROM delivery_notes WHERE id=? AND status='active'", (nid,)).fetchone()
    if not n:
        conn.close(); flash('Not found.', 'danger'); return redirect(url_for('notes'))
    items = conn.execute("SELECT * FROM dn_items WHERE dn_id=? ORDER BY sr_no", (nid,)).fetchall()
    conn.close()
    new_id = None
    conn = get_db()
    try:
        copy_cols = ("company_id,invoice_no,order_no,project_name,location,client_id,client_name,client_telephone,"
                     "invoice_address,shipping_address,show_prices,currency,vat_rate,subtotal,vat_amount,net_total,"
                     "vehicle_number,driver_name,driver_contact,notes,signatory_id,signatory_name,signatory_position,"
                     "signatory_phone,signatory_image,stamp_id,stamp_name,stamp_image,include_signature,include_stamp,"
                     "map_link,driver_id,task_type,pickup_name,pickup_phone,pickup_address,pickup_map_link").split(',')
        ttype = 'pickup' if is_pickup(n) else 'delivery'
        for _ in range(5):
            number = next_dn_number(n['company_id'] or 1, ttype)
            try:
                top = conn.execute("SELECT COALESCE(MAX(sort_order),0) AS m FROM delivery_notes", ()).fetchone()['m']
                conn.execute(returning_id(f"INSERT INTO delivery_notes(dn_number,dn_date,{','.join(copy_cols)},sort_order,created_by) "
                                          f"VALUES({','.join(['?'] * (len(copy_cols) + 4))})"),
                             [number, now_dubai().strftime('%Y-%m-%d')] + [n[c] for c in copy_cols] + [int(top) + 1, session['user_id']])
                new_id = last_insert_id(conn); break
            except Exception as e:
                try: conn._conn.rollback()
                except Exception: pass
                if 'unique' not in str(e).lower() and 'duplicate' not in str(e).lower(): raise
        for it in items:
            conn.execute("INSERT INTO dn_items(dn_id,sr_no,description,quantity,unit,amount,total_amount) VALUES(?,?,?,?,?,?,?)",
                         (new_id, it['sr_no'], it['description'], it['quantity'], it['unit'], it['amount'], it['total_amount']))
        conn.commit()
    finally:
        conn.close()
    audit('DN_DUPLICATED', f"{n['dn_number']} → {number}")
    flash(f'Duplicated as {number}. Review and adjust it below.', 'success')
    return redirect(url_for('edit_note', nid=new_id))

@app.route('/notes/<int:nid>/pdf')
@login_req
def download_pdf(nid):
    if not note_allowed(nid): return _no_access()
    conn = get_db(); n = conn.execute("SELECT * FROM delivery_notes WHERE id=?", (nid,)).fetchone(); conn.close()
    if not n: flash('Not found.', 'danger'); return redirect(url_for('notes'))
    p = n['pdf_path']
    if not p or not os.path.exists(p): p = gen_pdf(nid)
    if p and os.path.exists(p): return send_file(p, as_attachment=True, download_name=f"{n['dn_number']}.pdf")
    flash('PDF not available.', 'danger'); return redirect(url_for('view_note', nid=nid))

@app.route('/notes/<int:nid>/preview')
@login_req
def preview_pdf(nid):
    if not note_allowed(nid): return _no_access()
    conn = get_db(); n = conn.execute("SELECT * FROM delivery_notes WHERE id=?", (nid,)).fetchone(); conn.close()
    if not n: return "Not found", 404
    p = n['pdf_path']
    if not p or not os.path.exists(p): p = gen_pdf(nid)
    if p and os.path.exists(p): return send_file(p, mimetype='application/pdf')
    return "PDF generation failed", 500


# ── JSON APIs used by the form ─────────────────────────────────────────────
@app.route('/api/clients')
@login_req
def api_clients():
    q = request.args.get('q', '')
    conn = get_db()
    rows = conn.execute("SELECT * FROM clients WHERE is_active=1 AND (LOWER(name) LIKE LOWER(?) OR LOWER(COALESCE(company_name,'')) LIKE LOWER(?)) ORDER BY name LIMIT 20",
                        (f'%{q}%', f'%{q}%')).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])

@app.route('/api/clients/save', methods=['POST'])
@login_req
def api_client_save():
    """Called by the delivery form while typing: creates / updates the client straight away."""
    d = request.get_json(silent=True) or {}
    name = (d.get('name') or '').strip()
    if not name: return jsonify({'ok': False, 'error': 'name required'}), 400
    ml = d.get('map_link') or ''
    ml = clean_map_link(ml)[0] if ml else ''
    conn = get_db()
    try:
        cid = _autosave_client(conn, name, d.get('telephone'), d.get('address'), d.get('shipping_address'), ml,
                               client_id=_int_or_none(d.get('id')), last_used=bool(d.get('last_used', True)))
        conn.commit()
        row = conn.execute("SELECT * FROM clients WHERE id=?", (cid,)).fetchone()
    finally:
        conn.close()
    return jsonify({'ok': True, 'client': dict(row)})

@app.route('/api/stamps/<int:company_id>')
@login_req
def api_stamps(company_id):
    conn = get_db()
    rows = conn.execute("SELECT id,name,image_path,is_default FROM stamps WHERE is_active=1 AND company_id=? ORDER BY is_default DESC,name", (company_id,)).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])

@app.route('/api/next-number')
@login_req
def api_next_number():
    return jsonify({'number': next_dn_number(_int_or_none(request.args.get('co')) or 1, 'pickup' if request.args.get('type') == 'pickup' else 'delivery')})

@app.route('/api/drivers')
@login_req
def api_drivers():
    conn = get_db(); rows = conn.execute("SELECT * FROM drivers WHERE is_active=1 ORDER BY name").fetchall(); conn.close()
    return jsonify([dict(r) for r in rows])


# ── PDF GENERATION ─────────────────────────────────────────────────────────
def gen_pdf(nid):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, Spacer, Table, TableStyle, Image as RLImg,
                                    HRFlowable, KeepTogether, BaseDocTemplate, PageTemplate, Frame)
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT

    conn = get_db()
    n = conn.execute("SELECT * FROM delivery_notes WHERE id=?", (nid,)).fetchone()
    items = conn.execute("SELECT * FROM dn_items WHERE dn_id=? ORDER BY sr_no", (nid,)).fetchall()
    conn.close()
    if not n: return None
    co = get_company(n['company_id'])
    if not co: return None

    year = (n['dn_date'] or '')[:4] or str(now_dubai().year)
    ydir = os.path.join(ARCHIVE_DIR, year); os.makedirs(ydir, exist_ok=True)
    pdf_path = os.path.join(ydir, f"{n['dn_number']}.pdf")

    W, H = A4
    SM = 14 * mm; TM = 10 * mm
    FOOTER_H = 14 * mm; FOOTER_GAP = 5 * mm
    CW = W - 2 * SM
    content_bottom = FOOTER_H + FOOTER_GAP + 6 * mm

    BK = colors.black; WH = colors.white
    DK = colors.HexColor("#1a1a1a"); DK2 = colors.HexColor("#2c2c2c")
    TEAL = colors.HexColor("#111111")          # table header colour (black & white theme)
    BLUE = colors.HexColor("#000000")          # title colour
    GRID = colors.HexColor("#333333")
    BGRD = colors.HexColor("#cccccc")

    def P(txt, sz=8, bold=False, align=TA_LEFT, col=BK, raw=False, leading=None):
        t = str(txt or "") if raw else xml_escape(str(txt or "")).replace("\n", "<br/>")
        return Paragraph(t, ParagraphStyle("s", fontName="Helvetica-Bold" if bold else "Helvetica",
                         fontSize=sz, alignment=align, textColor=col,
                         leading=leading or sz * 1.3, spaceAfter=0, spaceBefore=0))

    from reportlab.lib.utils import ImageReader

    def small(path, max_px):
        """Down-scaled copy of an image so archived PDFs stay small (keeps transparency)."""
        from PIL import Image as PILImage
        im = PILImage.open(path); im.load()
        if max(im.size) > max_px:
            im.thumbnail((max_px, max_px), PILImage.LANCZOS)
        buf = io.BytesIO()
        if im.mode in ("RGBA", "LA", "P"):
            im.convert("RGBA").save(buf, "PNG", optimize=True)
        else:
            im.convert("RGB").save(buf, "JPEG", quality=88)
        buf.seek(0)
        return buf, im.size

    def fit_image(path, max_w, max_h):
        try:
            buf, (iw, ih) = small(path, 450)
            r = min(max_w / iw, max_h / ih)
            return RLImg(buf, width=iw * r, height=ih * r)
        except Exception:
            app.logger.exception('PDF image skipped: %s', path)
            return None

    def asset(rel):
        if not rel: return None
        p = os.path.join(UPLOAD_DIR, rel[len("uploads/"):]) if rel.startswith("uploads/") else os.path.join(BASE_DIR, "static", rel)
        return p if os.path.exists(p) else None

    logo_src = next((p for p in (os.path.join(UPLOAD_DIR, "logos", "GH_LOGO_pdf.jpg"),
                                 asset(co.get("logo_path")), os.path.join(UPLOAD_DIR, "logos", "GH_LOGO.jpg"))
                     if p and os.path.exists(p)), None)
    footer_src = os.path.join(IMAGES_DIR, "gh_footer_AE.png")
    show_prices = bool(n['show_prices'])
    try: footer_img = ImageReader(small(footer_src, 1400)[0]) if os.path.exists(footer_src) else None
    except Exception: footer_img = None

    def on_page(canv, doc):
        canv.saveState()
        if footer_img is not None:
            canv.drawImage(footer_img, SM, FOOTER_GAP, width=CW, height=FOOTER_H, preserveAspectRatio=False)
        canv.setFont("Helvetica", 7); canv.setFillColor(colors.HexColor("#777777"))
        canv.drawRightString(W - SM, FOOTER_GAP + FOOTER_H + 1.5 * mm, f"{n['dn_number']}  |  Page {doc.page}")
        canv.restoreState()

    elems = []

    # ── Letterhead ──
    LOGO_W = 22 * mm; GAP_W = 5 * mm; INFO_W = CW - LOGO_W - GAP_W
    addr_lines = [x for x in [co.get("address"), co.get("address2"), co.get("address3"), co.get("address4")] if x and x.strip()]
    contact = "   ".join(x for x in [("TEL: " + co["telephone"]) if co.get("telephone") else "",
                                    ("E-Mail: " + co["email"]) if co.get("email") else ""] if x)
    rows = [[P(co.get("name", ""), 11.5, True, TA_LEFT, DK)]] + [[P(a, 8, False, TA_LEFT, DK2)] for a in addr_lines]
    if contact: rows.append([P(contact, 8, False, TA_LEFT, DK2)])
    if co.get("website"): rows.append([P(co["website"], 8, False, TA_LEFT, DK2)])
    info = Table(rows, colWidths=[INFO_W])
    info.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 0.8), ("BOTTOMPADDING", (0, 0), (-1, -1), 0.8),
                              ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    logo_img = fit_image(logo_src, LOGO_W - 6, LOGO_W - 6) if logo_src else None
    lh = Table([[logo_img or P("GH", 16, True, TA_CENTER, WH), info]], colWidths=[LOGO_W + GAP_W, INFO_W])
    lh.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), DK), ("ALIGN", (0, 0), (0, -1), "CENTER"),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("TOPPADDING", (0, 0), (0, -1), 3), ("BOTTOMPADDING", (0, 0), (0, -1), 3),
                            ("LEFTPADDING", (0, 0), (0, -1), 3), ("RIGHTPADDING", (0, 0), (0, -1), 3),
                            ("LEFTPADDING", (1, 0), (1, -1), 2), ("TOPPADDING", (1, 0), (1, -1), 2),
                            ("BOTTOMPADDING", (1, 0), (1, -1), 2)]))
    elems += [lh, Spacer(1, 2 * mm), HRFlowable(width=CW, thickness=0.8, color=BGRD, spaceAfter=1 * mm)]

    # ── Title ──
    pickup_job = is_pickup(n)
    elems.append(P("<u>PICKUP &amp; DELIVERY NOTE</u>" if pickup_job else "<u>DELIVERY NOTE</u>", 16, True, TA_CENTER, BLUE, raw=True))
    elems.append(Spacer(1, 2 * mm))

    # ── Reference block ──
    def kv(label, value, bold_value=False):
        return P(f"<b>{xml_escape(label)}</b> {xml_escape(str(value or ''))}" if not bold_value
                 else f"<b>{xml_escape(label)} {xml_escape(str(value or ''))}</b>", 8.5, raw=True)
    left = [[kv("Pickup Note No :" if pickup_job else "Delivery Note No :", n['dn_number'])],
            [kv("Invoice No :", n['invoice_no'])],
            [kv("Order No :", n['order_no'])],
            [kv("Project name :", n['project_name'])],
            [kv("Date :", fmt_date_long(n['dn_date']))]]
    right = [[kv("Mr/Ms :", n['client_name'])],
             [kv("Tel No :", n['client_telephone'])],
             [kv("Location :", n['location'])],
             [P("", 8.5)], [P("", 8.5)]]
    ref = Table([[Table(left, colWidths=[CW * 0.56]), Table(right, colWidths=[CW * 0.44])]], colWidths=[CW * 0.56, CW * 0.44])
    inner_style = [("TOPPADDING", (0, 0), (-1, -1), 1.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2),
                   ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]
    for t in (ref._cellvalues[0][0], ref._cellvalues[0][1]): t.setStyle(TableStyle(inner_style))
    ref.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                             ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    elems += [ref, Spacer(1, 2.5 * mm)]

    # ── Addresses ──
    half = CW / 2
    if pickup_job:
        pu_txt = "\n".join(x for x in [n['pickup_name'], (("Tel: " + n['pickup_phone']) if n['pickup_phone'] else ""), n['pickup_address']] if x)
        left_h, left_t, right_h = "COLLECT FROM (PICKUP ADDRESS)", pu_txt, "DELIVER TO (DROP-OFF ADDRESS)"
        right_t = n['shipping_address'] or n['invoice_address']
    else:
        left_h, left_t, right_h, right_t = "INVOICE ADDRESS", n['invoice_address'], "SHIPPING ADDRESS", n['shipping_address']
    addr = Table([[P(left_h, 8.5, True, TA_CENTER, WH), P(right_h, 8.5, True, TA_CENTER, WH)],
                  [P(left_t, 8, True, TA_CENTER), P(right_t, 8, True, TA_CENTER)]],
                 colWidths=[half, half], rowHeights=[6.5 * mm, None])
    addr.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), TEAL), ("BOX", (0, 0), (-1, -1), 0.8, GRID),
                              ("INNERGRID", (0, 0), (-1, -1), 0.8, GRID), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                              ("TOPPADDING", (0, 1), (-1, 1), 4), ("BOTTOMPADDING", (0, 1), (-1, 1), 4)]))
    elems += [addr, Spacer(1, 2.5 * mm)]

    # ── Items ──
    if show_prices:
        cw = [11 * mm, None, 16 * mm, 16 * mm, 24 * mm, 28 * mm]
    else:
        cw = [12 * mm, None, 22 * mm, 24 * mm]
    cw[1] = CW - sum(c for c in cw if c)
    ncol = len(cw)
    hdr = [P("SL#", 7.5, False, TA_CENTER, WH), P("ITEM NAME", 8, True, TA_CENTER, WH),
           P("QTY", 7.5, False, TA_CENTER, WH), P("UNIT", 7.5, False, TA_CENTER, WH)]
    if show_prices: hdr += [P("AMOUNT", 7.5, False, TA_CENTER, WH), P("TOTAL AMOUNT", 7.5, False, TA_CENTER, WH)]
    data = [[P("DESCRIPTION OF GOODS", 9, True, TA_CENTER, WH)] + [""] * (ncol - 1), hdr]
    for it in items:
        row = [P(it['sr_no'], 8, False, TA_CENTER), P(it['description'], 8, True, TA_CENTER),
               P(fmt_qty(it['quantity']), 8.5, True, TA_CENTER), P(it['unit'], 8.5, True, TA_CENTER)]
        if show_prices:
            row += [P(f"{float(it['amount'] or 0):,.2f}", 8, False, TA_RIGHT), P(f"{float(it['total_amount'] or 0):,.2f}", 8, False, TA_RIGHT)]
        data.append(row)
    style = [("SPAN", (0, 0), (-1, 0)), ("BACKGROUND", (0, 0), (-1, 1), TEAL),
             ("GRID", (0, 0), (-1, -1), 0.7, GRID), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
             ("TOPPADDING", (0, 0), (-1, -1), 2.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6),
             ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
    if show_prices:
        cur = n['currency'] or 'AED'
        vr = float(n['vat_rate'] or 0); vr_txt = f"{vr:g}"
        def total_row(label, val, bold=False):
            r = [P(label, 8.5, bold, TA_RIGHT)] + [""] * (ncol - 2) + [P(val, 8.5, bold, TA_RIGHT)]
            data.append(r); i = len(data) - 1
            style.append(("SPAN", (0, i), (ncol - 2, i)))
        total_row("Total Amount :", f"{float(n['subtotal'] or 0):,.2f}")
        total_row(f"VAT {vr_txt}% :", f"{float(n['vat_amount'] or 0):,.2f}")
        total_row(f"Net payable amount ({cur}) :", f"{float(n['net_total'] or 0):,.2f}", True)
        data.append([P(f"Total Amount : {n2w(n['net_total'])}", 8.5)] + [""] * (ncol - 1))
        style.append(("SPAN", (0, len(data) - 1), (-1, len(data) - 1)))
    tbl = Table(data, colWidths=cw, repeatRows=2)
    tbl.setStyle(TableStyle(style))
    elems += [tbl, Spacer(1, 3 * mm)]

    # ── Vehicle / driver ──
    vw = [CW * 0.28, CW * 0.38, CW * 0.34]
    veh = Table([[P("VEHICLE NUMBER", 8.5, True, TA_CENTER, WH), P("DRIVER NAME", 8.5, True, TA_CENTER, WH),
                  P("DRIVER CONTACT NUMBER", 8.5, True, TA_CENTER, WH)],
                 [P(n['vehicle_number'], 9, True, TA_CENTER), P(n['driver_name'], 9, True, TA_CENTER),
                  P(n['driver_contact'], 9, True, TA_CENTER)]], colWidths=vw)
    veh.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), TEAL), ("GRID", (0, 0), (-1, -1), 0.8, GRID),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                             ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))

    # ── Signatures ──
    sw = CW / 4
    sig_t = Table([[P(x, 8, True, TA_CENTER, WH) for x in
                    ("SUPERVISOR SIGNATURE", "DISPATCH SIGNATURE", "DRIVER SIGNATURE", "RECEIVER SIGNATURE")],
                   ["", "", "", ""]], colWidths=[sw] * 4, rowHeights=[6 * mm, 15 * mm])
    sig_t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), TEAL), ("GRID", (0, 0), (-1, -1), 0.8, GRID),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))

    # ── Notes + receiver block ──
    note_paras = []
    for line in (n['notes'] or '').split("\n"):
        line = line.strip()
        if not line: continue
        m = re.match(r'^([^:]{1,40}):\s*(.*)$', line)
        txt = f"<b>{xml_escape(m.group(1))}:</b> {xml_escape(m.group(2))}" if m else xml_escape(line)
        note_paras.append(P(txt, 7, raw=True, leading=8.6))
        note_paras.append(Spacer(1, 1.6 * mm))
    rec = [P("<b>Name:..........................................</b>", 8, raw=True), Spacer(1, 3 * mm),
           P("<b>Tel no:........................................</b>", 8, raw=True), Spacer(1, 3 * mm),
           P("<b>Date:..........................................</b>", 8, raw=True)]
    nw = CW * 0.64
    notes_t = Table([[note_paras or [Spacer(1, 1)], rec]], colWidths=[nw, CW - nw])
    notes_t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                 ("RIGHTPADDING", (0, 0), (0, 0), 8), ("RIGHTPADDING", (1, 0), (1, 0), 0),
                                 ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                                 ("LEFTPADDING", (1, 0), (1, 0), 6), ("TOPPADDING", (1, 0), (1, 0), 3)]))

    # ── Thanks & regards: text | signature | stamp on one band ──
    sig_img = fit_image(asset(n['signatory_image']), 36 * mm, 15 * mm) if n['include_signature'] and n['signatory_image'] and asset(n['signatory_image']) else None
    stamp_img = fit_image(asset(n['stamp_image']), 28 * mm, 24 * mm) if n['include_stamp'] and n['stamp_image'] and asset(n['stamp_image']) else None
    left_col = [P("Thanks &amp; Regards", 8.5, raw=True), P(co.get("name", ""), 8.5), Spacer(1, 3.5 * mm)]
    if n['signatory_name']:
        left_col.append(P(("Mr. " if not n['signatory_name'].lower().startswith(("mr", "ms", "mrs", "dr")) else "") + n['signatory_name'], 8.5, True))
        if n['signatory_position']: left_col.append(P(n['signatory_position'], 8.5))
        if n['signatory_phone']: left_col.append(P(n['signatory_phone'], 8.5))
    regards = Table([[left_col, sig_img or "", stamp_img or ""]], colWidths=[CW * 0.46, CW * 0.28, CW * 0.26])
    regards.setStyle(TableStyle([("VALIGN", (0, 0), (0, 0), "TOP"), ("VALIGN", (1, 0), (1, 0), "BOTTOM"),
                                 ("VALIGN", (2, 0), (2, 0), "MIDDLE"), ("ALIGN", (1, 0), (2, 0), "CENTER"),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                 ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))

    elems.append(KeepTogether([veh, Spacer(1, 3 * mm), sig_t, Spacer(1, 2.5 * mm), notes_t, Spacer(1, 2 * mm), regards]))

    doc = BaseDocTemplate(pdf_path, pagesize=A4, leftMargin=SM, rightMargin=SM, topMargin=TM, bottomMargin=content_bottom,
                          title=n['dn_number'], author=co.get("name", ""))
    doc.addPageTemplates([PageTemplate(id="main", onPage=on_page,
                          frames=[Frame(SM, content_bottom, CW, H - TM - content_bottom,
                                        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)])])
    doc.build(elems)

    conn = get_db(); conn.execute("UPDATE delivery_notes SET pdf_path=? WHERE id=?", (pdf_path, nid)); conn.commit(); conn.close()
    return pdf_path


# ── DRIVER PORTAL ──────────────────────────────────────────────────────────
# Separate mobile-friendly portal. Each driver has their own login and only ever sees
# the delivery notes assigned to them (delivery_notes.driver_id = their driver record).
@app.route('/driver/login', methods=['GET', 'POST'])
def driver_login():
    if session.get('user_id') and session.get('role') == 'driver': return redirect(url_for('driver_home'))
    if request.method == 'POST':
        u = request.form.get('username', '').strip(); p = request.form.get('password', '')
        conn = get_db()
        user = conn.execute("SELECT * FROM users WHERE username=? AND is_active=1 AND role='driver'", (u,)).fetchone()
        drv = conn.execute("SELECT is_active FROM drivers WHERE id=?", (user['driver_id'],)).fetchone() if user and user['driver_id'] else None
        conn.close()
        if user and drv and drv['is_active'] and check_password_hash(user['password_hash'], p):
            session.clear()
            session['user_id'] = user['id']; session['username'] = user['username']
            session['full_name'] = user['full_name'] or user['username']; session['role'] = 'driver'
            session.permanent = True
            conn2 = get_db(); conn2.execute("UPDATE users SET last_login=CURRENT_TIMESTAMP WHERE id=?", (user['id'],)); conn2.commit(); conn2.close()
            audit('DRIVER_LOGIN'); return redirect(url_for('driver_home'))
        flash('Invalid username or password.', 'danger')
    return render_template('driver_login.html', s=gall())

@app.route('/driver/logout')
def driver_logout():
    audit('LOGOUT'); session.clear(); return redirect(url_for('driver_login'))

@app.route('/driver')
@driver_req
def driver_home():
    tab = request.args.get('tab', 'deliveries')
    if tab == 'pending': tab = 'deliveries'                      # old bookmarks
    if tab not in ('deliveries', 'pickups', 'done'): tab = 'deliveries'
    today = now_dubai().date()
    rng = request.args.get('range', '')
    d_from = request.args.get('from', '').strip(); d_to = request.args.get('to', '').strip()
    if rng == 'today': d_from = d_to = today.strftime('%Y-%m-%d')
    elif rng == 'tomorrow': d_from = d_to = (today + timedelta(days=1)).strftime('%Y-%m-%d')
    elif rng == 'week': d_from, d_to = today.strftime('%Y-%m-%d'), (today + timedelta(days=6)).strftime('%Y-%m-%d')
    elif rng == 'all': d_from = d_to = ''
    for v in (d_from, d_to):
        if v and not re.match(r'^\d{4}-\d{2}-\d{2}$', v): d_from = d_to = ''; break
    sort = request.args.get('sort', 'priority')
    if sort not in ('priority', 'date_asc', 'date_desc'): sort = 'priority'
    did = getattr(g, 'driver_id', None)
    conn = get_db()
    drv = conn.execute("SELECT * FROM drivers WHERE id=?", (did,)).fetchone() if did else None

    def cnt(extra, params=()):
        if not did: return 0
        return conn.execute("SELECT COUNT(*) c FROM delivery_notes WHERE status='active' AND driver_id=?" + extra, (did,) + tuple(params)).fetchone()['c']
    pend = " AND delivery_status='pending'"
    counts = {'deliveries': cnt(pend + " AND COALESCE(task_type,'delivery')='delivery'"),
              'pickups': cnt(pend + " AND task_type='pickup'"),
              'done': cnt(" AND delivery_status='delivered'")}

    sql = """SELECT n.*,c.code AS co_code,c.name AS co_name FROM delivery_notes n
             LEFT JOIN companies c ON n.company_id=c.id WHERE n.status='active' AND n.driver_id=?"""
    params = [did]
    if tab == 'done':
        sql += " AND n.delivery_status='delivered'"
        dcol = "COALESCE(n.delivered_on, n.dn_date)"
    else:
        sql += " AND n.delivery_status='pending' AND COALESCE(n.task_type,'delivery')=?"
        params.append('pickup' if tab == 'pickups' else 'delivery')
        dcol = "n.dn_date"
    if d_from: sql += f" AND {dcol}>=?"; params.append(d_from)
    if d_to: sql += f" AND {dcol}<=?"; params.append(d_to)
    if tab == 'done':
        order = {'priority': "COALESCE(n.delivered_at,'') DESC, n.id DESC",
                 'date_desc': f"{dcol} DESC, COALESCE(n.delivered_at,'') DESC",
                 'date_asc': f"{dcol} ASC, COALESCE(n.delivered_at,'') ASC"}[sort]
    else:
        order = {'priority': "n.dn_date, COALESCE(n.sort_order,0), n.id",
                 'date_asc': "n.dn_date ASC, COALESCE(n.sort_order,0), n.id",
                 'date_desc': "n.dn_date DESC, COALESCE(n.sort_order,0), n.id"}[sort]
    rows = conn.execute(sql + " ORDER BY " + order + " LIMIT 150", params).fetchall() if did else []
    items = {}
    if rows:
        ids = [r['id'] for r in rows]
        for it in conn.execute("SELECT * FROM dn_items WHERE dn_id IN (" + ','.join('?' * len(ids)) + ") ORDER BY sr_no", ids).fetchall():
            items.setdefault(it['dn_id'], []).append(it)
    conn.close()
    return render_template('driver_home.html', notes=rows, items=items, tab=tab, counts=counts, drv=drv, s=gall(),
                           d_from=d_from, d_to=d_to, rng=rng, sort=sort, today=today.strftime('%Y-%m-%d'),
                           tomorrow=(today + timedelta(days=1)).strftime('%Y-%m-%d'),
                           filtered=bool(d_from or d_to) or sort != 'priority')

def _back_to_driver_home():
    keep = {k: v for k, v in request.form.items() if k in ('tab', 'range', 'from', 'to', 'sort') and v}
    return redirect(url_for('driver_home', **keep))

@app.route('/driver/<int:nid>/pickup', methods=['POST'])
@driver_req
def driver_pickup(nid):
    """Step 1 of a pickup job: the driver confirms the goods were collected."""
    did = getattr(g, 'driver_id', None)
    conn = get_db()
    n = conn.execute("SELECT * FROM delivery_notes WHERE id=? AND driver_id=? AND status='active'", (nid, did)).fetchone() if did else None
    if not n or not is_pickup(n):
        conn.close(); flash('That pickup is not assigned to you.', 'danger'); return _back_to_driver_home()
    if n['delivery_status'] != 'pending' or n['picked_up_at']:
        conn.close(); flash(f"{n['dn_number']} is already collected.", 'danger'); return _back_to_driver_home()
    conn.execute("""UPDATE delivery_notes SET picked_up_at=?,picked_up_by=?,pickup_remarks=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                 (now_dubai().strftime('%Y-%m-%d %H:%M'), request.form.get('handed_by', '').strip(),
                  request.form.get('remarks', '').strip(), nid))
    conn.commit(); conn.close()
    audit('DN_PICKED_UP', f"{n['dn_number']} collected by driver")
    flash(f"{n['dn_number']} marked as collected. Now deliver it to the drop-off address.", 'success')
    return _back_to_driver_home()

@app.route('/driver/<int:nid>/deliver', methods=['POST'])
@driver_req
def driver_deliver(nid):
    did = getattr(g, 'driver_id', None)
    received_by = request.form.get('received_by', '').strip()
    conn = get_db()
    n = conn.execute("SELECT * FROM delivery_notes WHERE id=? AND driver_id=? AND status='active'", (nid, did)).fetchone() if did else None
    if not n:
        conn.close(); flash('That delivery is not assigned to you.', 'danger'); return _back_to_driver_home()
    if n['delivery_status'] != 'pending':
        conn.close(); flash(f"{n['dn_number']} is already {n['delivery_status']}.", 'danger'); return _back_to_driver_home()
    if is_pickup(n) and not n['picked_up_at']:
        conn.close(); flash('Mark the goods as collected first, then deliver them.', 'danger'); return _back_to_driver_home()
    if not received_by:
        conn.close(); flash('Enter the name of the person who received the goods.', 'danger'); return _back_to_driver_home()
    now = now_dubai()
    conn.execute("""UPDATE delivery_notes SET delivery_status='delivered',delivered_on=?,received_by=?,status_remarks=?,
                    delivered_by=?,delivered_at=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                 (now.strftime('%Y-%m-%d'), received_by, request.form.get('remarks', '').strip(),
                  f"Driver: {session.get('full_name') or session.get('username')}", now.strftime('%Y-%m-%d %H:%M'), nid))
    conn.commit(); conn.close()
    audit('DN_DELIVERED', f"{n['dn_number']} by driver, received by {received_by}")
    flash(f"{n['dn_number']} marked as delivered.", 'success')
    return _back_to_driver_home()

@app.route('/driver/password', methods=['GET', 'POST'])
@driver_req
def driver_password():
    if request.method == 'POST':
        cur = request.form.get('current_password', ''); nw = request.form.get('new_password', ''); cn = request.form.get('confirm_password', '')
        if nw != cn: flash('Passwords do not match.', 'danger')
        elif len(nw) < 6: flash('Minimum 6 characters.', 'danger')
        else:
            conn = get_db(); user = conn.execute("SELECT * FROM users WHERE id=?", (session['user_id'],)).fetchone()
            if user and check_password_hash(user['password_hash'], cur):
                conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(nw), session['user_id'])); conn.commit()
                flash('Password changed.', 'success')
            else: flash('Current password incorrect.', 'danger')
            conn.close()
    return render_template('driver_password.html', s=gall())


# ── COMPANIES ──────────────────────────────────────────────────────────────
@app.route('/companies')
@admin_req
def companies():
    return redirect(url_for('company_settings') + '#companies')

def _save_company_images(co=None):
    logo_path = co['logo_path'] if co else 'uploads/logos/GH_LOGO.jpg'
    stamp_path = co['stamp_path'] if co else ''
    for field, subdir in [('logo', 'logos'), ('stamp', 'stamps')]:
        fi = request.files.get(field)
        if fi and fi.filename and allowed_file(fi.filename):
            fn = secure_filename(fi.filename)
            fi.save(os.path.join(UPLOAD_DIR, subdir, fn))
            if field == 'logo': logo_path = f'uploads/{subdir}/{fn}'
            else: stamp_path = f'uploads/{subdir}/{fn}'
    return logo_path, stamp_path

@app.route('/companies/new', methods=['GET', 'POST'])
@admin_req
def new_company():
    if request.method == 'POST':
        f = request.form; conn = get_db()
        logo_path, stamp_path = _save_company_images()
        code = f.get('code', '').upper().strip()
        try:
            conn.execute("""INSERT INTO companies(code,name,address,address2,address3,address4,telephone,email,trn,website,
                logo_path,stamp_path,letterhead_path,dn_prefix,is_active,sort_order) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (code, f.get('name'), f.get('address'), f.get('address2', ''), f.get('address3', ''), f.get('address4', ''),
                 f.get('telephone'), f.get('email'), f.get('trn'), f.get('website'), logo_path, stamp_path, '',
                 f.get('dn_prefix') or f'{code}-DN', 1, int(f.get('sort_order', 99) or 99)))
            conn.commit(); audit('COMPANY_ADDED', code); flash('Company added.', 'success')
        except Exception as e:
            try: conn._conn.rollback()
            except Exception: pass
            flash(f'Could not add company: {e}', 'danger')
        finally:
            conn.close()
        return redirect(url_for('company_settings') + '#companies')
    return render_template('company_form.html', co=None, mode='new', s=gall())

@app.route('/companies/<int:cid>/edit', methods=['GET', 'POST'])
@admin_req
def edit_company(cid):
    conn = get_db(); co = conn.execute("SELECT * FROM companies WHERE id=?", (cid,)).fetchone()
    if not co:
        conn.close(); flash('Company not found.', 'danger'); return redirect(url_for('company_settings') + '#companies')
    if request.method == 'POST':
        f = request.form
        logo_path, stamp_path = _save_company_images(co)
        conn.execute("""UPDATE companies SET name=?,address=?,address2=?,address3=?,address4=?,telephone=?,email=?,trn=?,website=?,
            logo_path=?,stamp_path=?,dn_prefix=?,is_active=?,sort_order=? WHERE id=?""",
            (f.get('name'), f.get('address'), f.get('address2', ''), f.get('address3', ''), f.get('address4', ''),
             f.get('telephone'), f.get('email'), f.get('trn'), f.get('website'), logo_path, stamp_path,
             f.get('dn_prefix'), 1 if f.get('is_active') else 0, int(f.get('sort_order', 99) or 99), cid))
        conn.commit(); conn.close(); audit('COMPANY_EDITED', f.get('name')); flash('Company updated.', 'success')
        return redirect(url_for('company_settings') + '#companies')
    conn.close()
    return render_template('company_form.html', co=co, mode='edit', s=gall())

def _company_note_count(conn, cid):
    return conn.execute("SELECT COUNT(*) c FROM delivery_notes WHERE company_id=? AND status='active'", (cid,)).fetchone()['c']

@app.route('/companies/<int:cid>/delete', methods=['POST'])
@admin_req
def delete_company(cid):
    conn = get_db(); cnt = _company_note_count(conn, cid)
    if cnt > 0:
        flash(f'Cannot delete: {cnt} active delivery notes linked.', 'danger')
    else:
        conn.execute("DELETE FROM companies WHERE id=?", (cid,)); conn.commit()
        audit('COMPANY_DELETED', f'ID {cid}'); flash('Company deleted.', 'success')
    conn.close()
    return redirect(url_for('company_settings') + '#companies')

@app.route('/companies/bulk-delete', methods=['POST'])
@admin_req
def bulk_delete_companies():
    conn = get_db(); deleted = 0
    for cid in request.form.getlist('ids[]'):
        if _company_note_count(conn, cid) == 0:
            conn.execute("DELETE FROM companies WHERE id=?", (cid,)); deleted += 1
    conn.commit(); conn.close()
    flash(f'{deleted} compan{"y" if deleted == 1 else "ies"} deleted.', 'success')
    return redirect(url_for('company_settings') + '#companies')


# ── CLIENTS ────────────────────────────────────────────────────────────────
@app.route('/clients')
@login_req
def clients():
    q = request.args.get('q', ''); conn = get_db()
    if q:
        rows = conn.execute("SELECT * FROM clients WHERE is_active=1 AND (LOWER(name) LIKE LOWER(?) OR LOWER(COALESCE(company_name,'')) LIKE LOWER(?) OR LOWER(COALESCE(trn,'')) LIKE LOWER(?)) ORDER BY name",
                            (f'%{q}%', f'%{q}%', f'%{q}%')).fetchall()
    else:
        rows = conn.execute("SELECT * FROM clients WHERE is_active=1 ORDER BY name").fetchall()
    conn.close(); return render_template('clients.html', clients=rows, q=q, s=gall())

@app.route('/clients/new', methods=['GET', 'POST'])
@login_req
def new_client():
    if request.method == 'POST':
        f = request.form; conn = get_db()
        conn.execute("INSERT INTO clients(name,company_name,trn,address,country,telephone,email,notes,shipping_address,map_link) VALUES(?,?,?,?,?,?,?,?,?,?)",
                     (f.get('name'), f.get('company_name'), f.get('trn'), f.get('address'), f.get('country'), f.get('telephone'), f.get('email'), f.get('notes'),
                      f.get('shipping_address', '').strip(), clean_map_link(f.get('map_link'))[0]))
        conn.commit(); conn.close(); audit('CLIENT_ADDED', f.get('name')); flash('Client added.', 'success')
        return redirect(url_for('clients'))
    return render_template('client_form.html', cl=None, mode='new', s=gall())

@app.route('/clients/<int:cid>/edit', methods=['GET', 'POST'])
@login_req
def edit_client(cid):
    conn = get_db(); cl = conn.execute("SELECT * FROM clients WHERE id=?", (cid,)).fetchone()
    if request.method == 'POST':
        f = request.form
        conn.execute("UPDATE clients SET name=?,company_name=?,trn=?,address=?,country=?,telephone=?,email=?,notes=?,shipping_address=?,map_link=? WHERE id=?",
                     (f.get('name'), f.get('company_name'), f.get('trn'), f.get('address'), f.get('country'), f.get('telephone'), f.get('email'), f.get('notes'),
                      f.get('shipping_address', '').strip(), clean_map_link(f.get('map_link'))[0], cid))
        conn.commit(); conn.close(); audit('CLIENT_EDITED', f.get('name')); flash('Updated.', 'success')
        return redirect(url_for('clients'))
    conn.close(); return render_template('client_form.html', cl=cl, mode='edit', s=gall())

@app.route('/clients/<int:cid>/delete', methods=['POST'])
@admin_req
def delete_client(cid):
    conn = get_db(); conn.execute("UPDATE clients SET is_active=0 WHERE id=?", (cid,)); conn.commit(); conn.close()
    audit('CLIENT_DELETED', f'ID {cid}'); flash('Client archived.', 'success'); return redirect(url_for('clients'))


# ── DRIVERS ────────────────────────────────────────────────────────────────
@app.route('/drivers')
@login_req
def drivers():
    conn = get_db()
    rows = conn.execute("""SELECT d.*, u.username AS login_user, u.last_login AS login_last,
        (SELECT COUNT(*) FROM delivery_notes n WHERE n.driver_id=d.id AND n.status='active' AND n.delivery_status='pending') AS pending_count
        FROM drivers d LEFT JOIN users u ON u.driver_id=d.id AND u.role='driver'
        WHERE d.is_active=1 ORDER BY d.name""").fetchall()
    conn.close()
    return render_template('drivers.html', drivers=rows, s=gall())

@app.route('/drivers/save', methods=['POST'])
@login_req
def save_driver():
    f = request.form; did = _int_or_none(f.get('id')); name = f.get('name', '').strip()
    if not name:
        flash('Driver name is required.', 'danger'); return redirect(url_for('drivers'))
    conn = get_db()
    if did:
        conn.execute("UPDATE drivers SET name=?,contact=?,vehicle_number=? WHERE id=?",
                     (name, f.get('contact', '').strip(), f.get('vehicle_number', '').strip(), did))
    else:
        conn.execute("INSERT INTO drivers(name,contact,vehicle_number) VALUES(?,?,?)",
                     (name, f.get('contact', '').strip(), f.get('vehicle_number', '').strip()))
    conn.commit(); conn.close(); audit('DRIVER_SAVED', name); flash('Driver saved.', 'success')
    return redirect(url_for('drivers'))

@app.route('/drivers/<int:did>/delete', methods=['POST'])
@login_req
def delete_driver(did):
    conn = get_db(); conn.execute("UPDATE drivers SET is_active=0 WHERE id=?", (did,))
    conn.execute("UPDATE users SET is_active=0 WHERE driver_id=? AND role='driver'", (did,)); conn.commit(); conn.close()
    audit('DRIVER_REMOVED', f'ID {did}'); flash('Driver removed.', 'success'); return redirect(url_for('drivers'))


@app.route('/drivers/<int:did>/login', methods=['POST'])
@admin_req
def save_driver_login(did):
    """Create (or update / reset the password of) the driver's own portal login."""
    username = request.form.get('username', '').strip(); password = request.form.get('password', '')
    conn = get_db()
    try:
        drv = conn.execute("SELECT * FROM drivers WHERE id=? AND is_active=1", (did,)).fetchone()
        if not drv:
            flash('Driver not found.', 'danger'); return redirect(url_for('drivers'))
        ex = conn.execute("SELECT id FROM users WHERE driver_id=? AND role='driver'", (did,)).fetchone()
        if password and len(password) < 6:
            flash('Password must be at least 6 characters.', 'danger'); return redirect(url_for('drivers'))
        if ex:
            if username: conn.execute("UPDATE users SET username=?,is_active=1 WHERE id=?", (username, ex['id']))
            else: conn.execute("UPDATE users SET is_active=1 WHERE id=?", (ex['id'],))
            if password: conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(password), ex['id']))
            msg = 'Driver login updated.'
        else:
            if not username or not password:
                flash('Enter a username and a password (min 6 characters) for the driver login.', 'danger'); return redirect(url_for('drivers'))
            conn.execute("INSERT INTO users(username,password_hash,full_name,role,permissions,driver_id) VALUES(?,?,?,?,?,?)",
                         (username, generate_password_hash(password), drv['name'], 'driver', '', did))
            msg = f"Driver login created. {drv['name']} signs in at /driver/login."
        conn.commit(); audit('DRIVER_LOGIN_SAVED', f"{drv['name']} ({username or 'password reset'})"); flash(msg, 'success')
    except Exception as e:
        try: conn._conn.rollback()
        except Exception: pass
        flash('Could not save the login - that username may already be taken.', 'danger')
    finally:
        conn.close()
    return redirect(url_for('drivers'))

@app.route('/drivers/<int:did>/login/delete', methods=['POST'])
@admin_req
def delete_driver_login(did):
    conn = get_db(); conn.execute("DELETE FROM users WHERE driver_id=? AND role='driver'", (did,)); conn.commit(); conn.close()
    audit('DRIVER_LOGIN_REMOVED', f'driver ID {did}'); flash('Driver login removed.', 'success'); return redirect(url_for('drivers'))


# ── SIGNATORIES & STAMPS ───────────────────────────────────────────────────
@app.route('/signatories')
@admin_req
def signatories():
    conn = get_db()
    sigs = conn.execute("SELECT * FROM signatories ORDER BY is_default DESC,name").fetchall()
    stamps = conn.execute("SELECT s.*,c.code as co_code,c.name as co_name FROM stamps s LEFT JOIN companies c ON s.company_id=c.id ORDER BY c.sort_order,s.is_default DESC").fetchall()
    companies = conn.execute("SELECT * FROM companies WHERE is_active=1 ORDER BY sort_order").fetchall()
    conn.close(); return render_template('signatories.html', sigs=sigs, stamps=stamps, companies=companies, s=gall())

@app.route('/signatories/add', methods=['POST'])
@admin_req
def add_signatory():
    name = request.form.get('name', ''); pos = request.form.get('position', ''); phone = request.form.get('phone', '')
    isdef = 1 if request.form.get('is_default') else 0; imgp = ''
    fi = request.files.get('image')
    if fi and fi.filename and allowed_file(fi.filename):
        fn = secure_filename(fi.filename); fi.save(os.path.join(UPLOAD_DIR, 'signatures', fn))
        imgp = f'uploads/signatures/{fn}'
    conn = get_db()
    if isdef: conn.execute("UPDATE signatories SET is_default=0")
    conn.execute("INSERT INTO signatories(name,position,phone,image_path,is_default) VALUES(?,?,?,?,?)", (name, pos, phone, imgp, isdef))
    conn.commit(); conn.close(); audit('SIGNATORY_ADDED', name); flash('Signatory added.', 'success')
    return redirect(url_for('signatories'))

@app.route('/signatories/<int:sid>/delete', methods=['POST'])
@admin_req
def delete_signatory(sid):
    conn = get_db(); conn.execute("DELETE FROM signatories WHERE id=?", (sid,)); conn.commit(); conn.close()
    flash('Signatory deleted.', 'success'); return redirect(url_for('signatories'))

@app.route('/signatories/<int:sid>/toggle', methods=['POST'])
@admin_req
def toggle_signatory(sid):
    conn = get_db(); cur = conn.execute("SELECT is_active FROM signatories WHERE id=?", (sid,)).fetchone()
    if cur: conn.execute("UPDATE signatories SET is_active=? WHERE id=?", (0 if cur['is_active'] else 1, sid))
    conn.commit(); conn.close(); flash('Updated.', 'success'); return redirect(url_for('signatories'))

@app.route('/signatories/<int:sid>/set-default', methods=['POST'])
@admin_req
def set_default_sig(sid):
    conn = get_db(); conn.execute("UPDATE signatories SET is_default=0"); conn.execute("UPDATE signatories SET is_default=1 WHERE id=?", (sid,)); conn.commit(); conn.close()
    flash('Default set.', 'success'); return redirect(url_for('signatories'))

@app.route('/stamps/add', methods=['POST'])
@admin_req
def add_stamp():
    name = request.form.get('name', ''); co_id = int(request.form.get('company_id', 1))
    isdef = 1 if request.form.get('is_default') else 0; imgp = ''
    fi = request.files.get('image')
    if fi and fi.filename and allowed_file(fi.filename):
        fn = secure_filename(fi.filename); fi.save(os.path.join(UPLOAD_DIR, 'stamps', fn))
        imgp = f'uploads/stamps/{fn}'
    conn = get_db()
    if isdef: conn.execute("UPDATE stamps SET is_default=0 WHERE company_id=?", (co_id,))
    conn.execute("INSERT INTO stamps(name,company_id,image_path,is_default) VALUES(?,?,?,?)", (name, co_id, imgp, isdef))
    conn.commit(); conn.close(); audit('STAMP_ADDED', name); flash('Stamp added.', 'success')
    return redirect(url_for('signatories'))

@app.route('/stamps/<int:sid>/delete', methods=['POST'])
@admin_req
def delete_stamp(sid):
    conn = get_db(); conn.execute("DELETE FROM stamps WHERE id=?", (sid,)); conn.commit(); conn.close()
    flash('Stamp deleted.', 'success'); return redirect(url_for('signatories'))

@app.route('/stamps/<int:sid>/toggle', methods=['POST'])
@admin_req
def toggle_stamp(sid):
    conn = get_db(); cur = conn.execute("SELECT is_active FROM stamps WHERE id=?", (sid,)).fetchone()
    if cur: conn.execute("UPDATE stamps SET is_active=? WHERE id=?", (0 if cur['is_active'] else 1, sid))
    conn.commit(); conn.close(); flash('Updated.', 'success'); return redirect(url_for('signatories'))

@app.route('/stamps/<int:sid>/set-default', methods=['POST'])
@admin_req
def set_default_stamp(sid):
    conn = get_db(); stmp = conn.execute("SELECT company_id FROM stamps WHERE id=?", (sid,)).fetchone()
    if stmp:
        conn.execute("UPDATE stamps SET is_default=0 WHERE company_id=?", (stmp['company_id'],))
        conn.execute("UPDATE stamps SET is_default=1 WHERE id=?", (sid,))
    conn.commit(); conn.close(); flash('Default stamp set.', 'success'); return redirect(url_for('signatories'))


# ── USERS ──────────────────────────────────────────────────────────────────
@app.route('/users')
@admin_req
def users():
    conn = get_db(); rows = conn.execute("SELECT * FROM users ORDER BY role,username").fetchall(); conn.close()
    return render_template('users.html', users=rows, s=gall(), perm_defs=DN_PERMS, user_perms=user_perms)

@app.route('/users/new', methods=['GET', 'POST'])
@admin_req
def new_user():
    if request.method == 'POST':
        f = request.form; conn = get_db()
        try:
            role = f.get('role', 'user') if f.get('role') in ('admin', 'user') else 'user'
            conn.execute("INSERT INTO users(username,password_hash,full_name,email,role,permissions) VALUES(?,?,?,?,?,?)",
                         (f.get('username'), generate_password_hash(f.get('password', '')), f.get('full_name'), f.get('email'), role,
                          ','.join(p for p in DN_PERMS if p in f.getlist('perm'))))
            conn.commit(); audit('USER_CREATED', f.get('username')); flash('User created.', 'success')
        except Exception as e:
            try: conn._conn.rollback()
            except Exception: pass
            flash(f'Error: {e}', 'danger')
        finally: conn.close()
        return redirect(url_for('users'))
    return render_template('user_form.html', u=None, mode='new', s=gall(), perm_defs=DN_PERMS, u_perms=set(DEFAULT_PERMS))

@app.route('/users/<int:uid>/edit', methods=['GET', 'POST'])
@admin_req
def edit_user(uid):
    conn = get_db(); u = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    if u and u['role'] == 'driver':
        conn.close(); flash('Driver logins are managed on the Drivers page.', 'danger'); return redirect(url_for('drivers'))
    if request.method == 'POST':
        f = request.form
        role = f.get('role') if f.get('role') in ('admin', 'user') else 'user'
        active = 1 if f.get('is_active') else 0
        if uid == session.get('user_id'): role = 'admin'; active = 1      # never lock yourself out
        conn.execute("UPDATE users SET full_name=?,email=?,role=?,is_active=?,permissions=? WHERE id=?",
                     (f.get('full_name'), f.get('email'), role, active, ','.join(p for p in DN_PERMS if p in f.getlist('perm')), uid))
        if f.get('password'): conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(f.get('password')), uid))
        conn.commit(); conn.close(); audit('USER_EDITED', f.get('full_name')); flash('Updated.', 'success')
        return redirect(url_for('users'))
    conn.close()
    return render_template('user_form.html', u=u, mode='edit', s=gall(), perm_defs=DN_PERMS, u_perms=user_perms(u['role'], u['permissions']))

@app.route('/users/<int:uid>/delete', methods=['POST'])
@admin_req
def delete_user(uid):
    if uid == session.get('user_id'): flash('Cannot delete yourself.', 'danger'); return redirect(url_for('users'))
    conn = get_db(); conn.execute("DELETE FROM users WHERE id=?", (uid,)); conn.commit(); conn.close()
    audit('USER_DELETED', f'ID {uid}'); flash('User deleted.', 'success'); return redirect(url_for('users'))


# ── SETTINGS ───────────────────────────────────────────────────────────────
@app.route('/settings', methods=['GET', 'POST'])
@admin_req
def company_settings():
    if request.method == 'POST':
        f = request.form
        for k in ['default_vat_rate', 'default_company_id', 'default_dn_notes']:
            if k in f: sset(k, f.get(k))
        fi = request.files.get('logo')
        if fi and fi.filename and allowed_file(fi.filename):
            fn = secure_filename(fi.filename); fi.save(os.path.join(UPLOAD_DIR, 'logos', fn))
            sset('logo_path', f'uploads/logos/{fn}')
        audit('SETTINGS_UPDATED'); flash('Settings saved.', 'success')
        return redirect(url_for('company_settings'))
    conn = get_db(); all_companies = conn.execute("SELECT * FROM companies ORDER BY sort_order,code").fetchall(); conn.close()
    return render_template('settings.html', s=gall(), companies=get_companies(), all_companies=all_companies)

@app.route('/audit-log')
@admin_req
def audit_log():
    conn = get_db(); logs = conn.execute("SELECT * FROM audit_logs ORDER BY id DESC LIMIT 500").fetchall(); conn.close()
    return render_template('audit_log.html', logs=logs, s=gall())


# ── ARCHIVE ────────────────────────────────────────────────────────────────
@app.route('/archive')
@admin_req
def archive():
    years = []
    if os.path.exists(ARCHIVE_DIR):
        for y in sorted(os.listdir(ARCHIVE_DIR), reverse=True):
            yd = os.path.join(ARCHIVE_DIR, y)
            if os.path.isdir(yd):
                fls = sorted(os.listdir(yd), reverse=True)
                years.append({'year': y, 'files': fls, 'count': len(fls)})
    return render_template('archive.html', years=years, archive_dir=ARCHIVE_DIR, s=gall())

@app.route('/archive/<year>/<fn>')
@admin_req
def dl_archive(year, fn):
    return send_from_directory(os.path.join(ARCHIVE_DIR, year), fn, as_attachment=True)


# ── BACKUP / RESTORE ───────────────────────────────────────────────────────
BACKUP_TABLES = ['users', 'settings', 'companies', 'clients', 'signatories', 'stamps', 'drivers',
                 'delivery_notes', 'dn_items', 'audit_logs']

def _table_cols(conn, t):
    """[(column, type)] for a table, on SQLite and PostgreSQL."""
    if USE_PG:
        rows = conn.execute("SELECT column_name AS name, data_type AS type FROM information_schema.columns "
                            "WHERE table_schema=current_schema() AND table_name=? ORDER BY ordinal_position", (t,)).fetchall()
        return [(r['name'], (r['type'] or '').lower()) for r in rows]
    return [(r['name'], (r['type'] or '').lower()) for r in conn.execute(f"PRAGMA table_info({t})").fetchall()]

_ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

@app.route('/backup')
@admin_req
def backup():
    """One-click full backup: every table + uploaded logos/stamps/signatures in a single Excel file."""
    from openpyxl import Workbook
    ts = now_dubai().strftime('%Y%m%d_%H%M%S')
    wb = Workbook(); meta = wb.active; meta.title = '_meta'
    conn = get_db(); counts = {}
    for t in BACKUP_TABLES:
        cols = [c for c,_ in _table_cols(conn, t)]
        rows = conn.execute(f"SELECT * FROM {t}").fetchall()
        ws = wb.create_sheet(t); ws.append(cols)
        for r in rows:
            ws.append([None]*len(cols))
            for ci, c in enumerate(cols, 1):
                v = r[c]
                if isinstance(v, str): v = _ILLEGAL.sub('', v)
                cell = ws.cell(row=ws.max_row, column=ci, value=v)
                if isinstance(v, str): cell.data_type = 's'   # never treat text such as "=abc" as a formula
        counts[t] = len(rows)
    conn.close()
    fs = wb.create_sheet('_files'); fs.append(['path','chunk','data'])
    nfiles = 0
    for root, dirs, files in os.walk(UPLOAD_DIR):
        for fn in sorted(files):
            fp = os.path.join(root, fn); rel = 'uploads/' + os.path.relpath(fp, UPLOAD_DIR).replace(os.sep, '/')
            with open(fp, 'rb') as fh: b64 = base64.b64encode(fh.read()).decode()
            for i in range(0, max(len(b64),1), 30000):
                fs.append([rel, i//30000, b64[i:i+30000]])
            nfiles += 1
    meta.append(['key','value'])
    for k, v in [('app','GH Delivery Backup'),('format_version',1),('created',now_dubai().strftime('%Y-%m-%d %H:%M:%S')),
                 ('database','PostgreSQL' if USE_PG else 'SQLite'),('files',nfiles)] + [(f'rows:{t}',n) for t,n in counts.items()]:
        meta.append([k, v])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    audit('BACKUP_CREATED', f'GH_Backup_{ts}.xlsx')
    return send_file(buf, as_attachment=True, download_name=f'GH_Backup_{ts}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ── shared restore engine (used by Excel and full-ZIP restore) ─────────────
def _coerce_for_column(v, ty):
    if v is None: return None
    if 'int' in ty and isinstance(v, float) and v.is_integer(): return int(v)
    if ('char' in ty or 'text' in ty) and not isinstance(v, str):
        return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
    if ty in ('real', 'double precision', 'numeric', 'float') and isinstance(v, int) and not isinstance(v, bool): return float(v)
    return v

def _read_sequences(conn):
    """Last id handed out per table, so ids continue exactly as before after a restore."""
    seqs = {}
    try:
        if USE_PG:
            for t in BACKUP_TABLES:
                if 'id' not in dict(_table_cols(conn, t)): continue
                sn = conn.execute("SELECT pg_get_serial_sequence(?, 'id') AS sn", (t,)).fetchone()['sn']
                if not sn: continue
                r = conn.execute(f"SELECT last_value, is_called FROM {sn}").fetchone()
                seqs[t] = int(r['last_value']) if r['is_called'] else int(r['last_value']) - 1
        else:
            for r in conn.execute("SELECT name, seq FROM sqlite_sequence").fetchall():
                if r['name'] in BACKUP_TABLES: seqs[r['name']] = int(r['seq'])
    except Exception:
        app.logger.exception('could not read id sequences')
    return seqs

def _replace_all_tables(conn, tables, sequences=None):
    """Replace the content of every backup table. tables = {name: (columns, rows)}. Caller commits/rolls back."""
    info = {t: _table_cols(conn, t) for t in BACKUP_TABLES}
    for t in reversed(BACKUP_TABLES): conn.execute(f"DELETE FROM {t}")
    counts = {}
    for t in BACKUP_TABLES:
        header, rows = tables[t]
        types = dict(info[t])
        use = [(i, h) for i, h in enumerate(header) if h in types]
        n = 0
        for r in rows:
            if all(v is None for v in r): continue
            vals = [_coerce_for_column(r[i] if i < len(r) else None, types[h]) for i, h in use]
            conn.execute(f"INSERT INTO {t}({','.join(h for _,h in use)}) VALUES({','.join(['?']*len(use))})", vals)
            n += 1
        counts[t] = n
        if 'id' in types:
            maxid = conn.execute(f"SELECT COALESCE(MAX(id),0) AS m FROM {t}").fetchone()['m']
            last = max(int((sequences or {}).get(t, 0)), int(maxid))
            if USE_PG:
                conn.execute(f"SELECT setval(pg_get_serial_sequence('{t}','id'), ?, ?)", (max(last, 1), last >= 1))
            elif sequences and t in sequences:
                try:
                    if conn.execute("SELECT 1 FROM sqlite_sequence WHERE name=?", (t,)).fetchone():
                        conn.execute("UPDATE sqlite_sequence SET seq=? WHERE name=?", (last, t))
                    else:
                        conn.execute("INSERT INTO sqlite_sequence(name,seq) VALUES(?,?)", (t, last))
                except Exception:
                    app.logger.exception('could not restore sqlite sequence for %s', t)
    return counts

def _ensure_upload_dirs():
    for d in (UPLOAD_DIR, ARCHIVE_DIR):
        os.makedirs(d, exist_ok=True)
    for sub in ('logos', 'stamps', 'signatures'):
        os.makedirs(os.path.join(UPLOAD_DIR, sub), exist_ok=True)

# ── FULL ZIP BACKUP ────────────────────────────────────────────────────────
FULL_BACKUP_APP = 'GH Delivery Full Backup'
_FILE_ROOTS = (('uploads', lambda: UPLOAD_DIR), ('archive', lambda: ARCHIVE_DIR))

def _json_default(o):
    from decimal import Decimal
    if isinstance(o, Decimal): return float(o)
    if hasattr(o, 'isoformat'): return o.isoformat(sep=' ')
    return str(o)

def _sha256_bytes(b):
    import hashlib; return hashlib.sha256(b).hexdigest()

def _pdf_path_out(v):
    """Store archive PDF paths relative to the archive folder so they work on any server."""
    if isinstance(v, str) and v:
        root = os.path.normpath(ARCHIVE_DIR) + os.sep
        if os.path.normpath(v).startswith(root):
            return '@archive/' + os.path.relpath(os.path.normpath(v), ARCHIVE_DIR).replace(os.sep, '/')
    return v

def _pdf_path_in(v):
    if isinstance(v, str) and v.startswith('@archive/'):
        return os.path.join(ARCHIVE_DIR, *v[len('@archive/'):].split('/'))
    return v

def _build_full_backup(fobj):
    """Write the complete backup (all tables exactly + every uploaded/archived file + checksums) into fobj."""
    import zipfile
    manifest = {'app': FULL_BACKUP_APP, 'format_version': 1,
                'created': now_dubai().strftime('%Y-%m-%d %H:%M:%S'),
                'database': 'PostgreSQL' if USE_PG else 'SQLite', 'tables': {}, 'sequences': {}, 'files': []}
    with zipfile.ZipFile(fobj, 'w', zipfile.ZIP_DEFLATED) as z:
        conn = get_db()
        for t in BACKUP_TABLES:
            cols = [c for c, _ in _table_cols(conn, t)]
            rows = conn.execute(f"SELECT * FROM {t}").fetchall()
            data = []
            for r in rows:
                data.append([_pdf_path_out(r[c]) if (t == 'delivery_notes' and c == 'pdf_path') else r[c] for c in cols])
            payload = json.dumps({'table': t, 'columns': cols, 'rows': data}, ensure_ascii=False, default=_json_default).encode('utf-8')
            z.writestr(f'database/{t}.json', payload)
            manifest['tables'][t] = {'rows': len(rows), 'sha256': _sha256_bytes(payload)}
        manifest['sequences'] = _read_sequences(conn)
        conn.close()
        for label, getroot in _FILE_ROOTS:
            root = getroot()
            for dp, dn, fns in os.walk(root):
                dn.sort()
                for fn in sorted(fns):
                    fp = os.path.join(dp, fn)
                    if fn.endswith('.restore_tmp'): continue
                    rel = os.path.relpath(fp, root).replace(os.sep, '/')
                    arc = f'files/{label}/{rel}'
                    with open(fp, 'rb') as fh: data = fh.read()
                    mt = os.path.getmtime(fp)
                    dt = time.localtime(mt)[:6]
                    zi = zipfile.ZipInfo(arc, date_time=dt if dt[0] >= 1980 else (1980, 1, 1, 0, 0, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    z.writestr(zi, data)
                    manifest['files'].append({'path': arc, 'size': len(data), 'sha256': _sha256_bytes(data), 'mtime': mt})
        z.writestr('manifest.json', json.dumps(manifest, indent=1))
        z.writestr('README.txt',
            "GH Delivery - full backup\n"
            f"Created: {manifest['created']} (Dubai time)\n\n"
            "database/*.json   every table, all columns, exactly as stored (ids included)\n"
            "files/uploads/    logos, stamps, signatures, letterheads\n"
            "files/archive/    generated delivery note PDFs\n"
            "manifest.json     row counts and SHA-256 checksums of everything above\n\n"
            "To restore: sign in as admin > Restore Backup > choose this .zip file.\n"
            "Restoring replaces ALL current data with the contents of this file.\n")
    return manifest

def _save_prestore_snapshot():
    """Safety copy of the current state, kept on the server in ./Backups (newest 3)."""
    try:
        d = BACKUP_DIR; os.makedirs(d, exist_ok=True)
        name = f"pre-restore_{now_dubai().strftime('%Y%m%d_%H%M%S')}.zip"
        with open(os.path.join(d, name), 'wb') as fh: _build_full_backup(fh)
        old = sorted(x for x in os.listdir(d) if x.startswith('pre-restore_') and x.endswith('.zip'))
        for x in old[:-3]:
            try: os.remove(os.path.join(d, x))
            except OSError: pass
        return name
    except Exception:
        app.logger.exception('pre-restore snapshot failed')
        return None

@app.route('/backup/zip')
@admin_req
def backup_zip():
    """One-click complete backup (database + all files) as a single ZIP."""
    import tempfile
    tmp = tempfile.TemporaryFile()
    man = _build_full_backup(tmp); tmp.seek(0)
    audit('FULL_BACKUP_CREATED', f"{sum(t['rows'] for t in man['tables'].values())} rows, {len(man['files'])} files")
    return send_file(tmp, as_attachment=True, download_name=f"GH_FullBackup_{now_dubai().strftime('%Y%m%d_%H%M%S')}.zip",
                     mimetype='application/zip')

def _restore_full_zip(fi):
    """Verify first (nothing is touched unless every checksum matches), then restore database + files exactly."""
    import zipfile, hashlib, shutil
    try:
        z = zipfile.ZipFile(fi.stream)
        bad = z.testzip()
        if bad: raise ValueError(f'Corrupt file inside the zip: {bad}')
        man = json.loads(z.read('manifest.json'))
        if man.get('app') != FULL_BACKUP_APP: raise ValueError('This is not a GH Delivery full backup.')
        tables = {}
        for t in BACKUP_TABLES:
            meta = (man.get('tables') or {}).get(t)
            if meta is None: raise ValueError(f'Table "{t}" is missing from the backup.')
            raw = z.read(f'database/{t}.json')
            if _sha256_bytes(raw) != meta['sha256']: raise ValueError(f'Checksum mismatch for table "{t}" - the backup is damaged.')
            d = json.loads(raw)
            if len(d['rows']) != meta['rows']: raise ValueError(f'Row count mismatch for table "{t}".')
            cols = d['columns']
            if t == 'delivery_notes' and 'pdf_path' in cols:
                pi = cols.index('pdf_path')
                for r in d['rows']: r[pi] = _pdf_path_in(r[pi])
            tables[t] = (cols, d['rows'])
        entries = {}
        for e in man.get('files', []):
            p = e['path']; parts = p.split('/')
            if len(parts) < 3 or parts[0] != 'files' or parts[1] not in ('uploads', 'archive') or '..' in parts or p.startswith('/') or '\\' in p:
                raise ValueError(f'Unsafe path in backup: {p}')
            h = hashlib.sha256(); size = 0
            with z.open(p) as src:
                for chunk in iter(lambda: src.read(1 << 20), b''):
                    h.update(chunk); size += len(chunk)
            if h.hexdigest() != e['sha256'] or size != e['size']: raise ValueError(f'Checksum mismatch for file {p}.')
            entries[p] = e
    except Exception as ex:
        flash(f'Cannot restore: {ex}', 'danger'); return redirect(url_for('restore_backup'))

    snapshot = _save_prestore_snapshot()
    conn = get_db()
    try:
        counts = _replace_all_tables(conn, tables, man.get('sequences'))
        conn.commit()
    except Exception as ex:
        try: conn._conn.rollback()
        except Exception: pass
        conn.close(); app.logger.exception('zip restore failed (database)')
        flash(f'Restore failed, nothing was changed: {ex}', 'danger'); return redirect(url_for('restore_backup'))
    conn.close()

    try:
        for label, getroot in _FILE_ROOTS:
            root = getroot(); tmp = root + '.restore_tmp'; old = root + '.restore_old'
            for p in (tmp, old):
                if os.path.exists(p): shutil.rmtree(p, ignore_errors=True)
            os.makedirs(tmp)
            prefix = f'files/{label}/'
            for p, e in entries.items():
                if not p.startswith(prefix): continue
                dest = os.path.join(tmp, *p[len(prefix):].split('/'))
                os.makedirs(os.path.dirname(dest), exist_ok=True)
                with z.open(p) as src, open(dest, 'wb') as out: shutil.copyfileobj(src, out)
                if e.get('mtime'): os.utime(dest, (e['mtime'], e['mtime']))
            if os.path.exists(root): os.replace(root, old)
            os.replace(tmp, root)
            shutil.rmtree(old, ignore_errors=True)
        _ensure_upload_dirs()
    except Exception as ex:
        app.logger.exception('zip restore failed (files)')
        flash(f'Database restored, but writing files failed: {ex}. A safety copy is in the Backups folder.', 'danger')
        return redirect(url_for('login'))
    app.logger.info('Full backup restored: %s', counts)
    session.clear()
    flash(f"Full backup restored exactly ({counts.get('delivery_notes',0)} delivery notes, {counts.get('companies',0)} companies, "
          f"{len(entries)} files). Please log in again." + (f" A safety copy of the previous data was saved as Backups/{snapshot}." if snapshot else ''), 'success')
    return redirect(url_for('login'))

@app.route('/restore', methods=['GET','POST'])
@admin_req
def restore_backup():
    if request.method == 'GET':
        return render_template('restore.html', s=gall())
    fi = request.files.get('backup_file')
    name = (fi.filename or '').lower() if fi else ''
    if name.endswith('.zip'): return _restore_full_zip(fi)
    from openpyxl import load_workbook
    if not fi or not name.endswith('.xlsx'):
        flash('Please choose a GH backup file (.zip full backup or .xlsx).', 'danger'); return redirect(url_for('restore_backup'))
    try:
        wb = load_workbook(fi, data_only=True)
        if '_meta' not in wb.sheetnames or any(t not in wb.sheetnames for t in BACKUP_TABLES):
            raise ValueError('This is not a valid GH Delivery backup file.')
        mk = {r[0]: r[1] for r in wb['_meta'].iter_rows(min_row=2, values_only=True) if r and r[0]}
        if mk.get('app') != 'GH Delivery Backup': raise ValueError('This is not a valid GH Delivery backup file.')
    except Exception as e:
        flash(f'Cannot read backup: {e}', 'danger'); return redirect(url_for('restore_backup'))

    conn = get_db()
    try:
        tables = {}
        for t in BACKUP_TABLES:
            rows = wb[t].iter_rows(values_only=True)
            header = list(next(rows, None) or ())
            tables[t] = (header, rows)
        restored = _replace_all_tables(conn, tables)
        conn.commit()
    except Exception as e:
        try: conn._conn.rollback()
        except Exception: pass
        conn.close(); app.logger.exception('restore failed')
        flash(f'Restore failed, nothing was changed: {e}', 'danger'); return redirect(url_for('restore_backup'))
    conn.close()

    # restore uploaded images (logos, stamps, signatures)
    chunks = {}
    for r in wb['_files'].iter_rows(min_row=2, values_only=True):
        if r and r[0]: chunks.setdefault(r[0], []).append((int(r[1] or 0), r[2] or ''))
    nf = 0; up_root = os.path.abspath(UPLOAD_DIR)
    for rel, parts in chunks.items():
        dest = os.path.abspath(os.path.join(UPLOAD_DIR, rel[len('uploads/'):] if rel.startswith('uploads/') else rel))
        if not dest.startswith(up_root + os.sep): continue
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        with open(dest, 'wb') as fh: fh.write(base64.b64decode(''.join(d for _, d in sorted(parts))))
        nf += 1
    audit('BACKUP_RESTORED', f"{fi.filename}: " + ', '.join(f'{t}={n}' for t, n in restored.items()))
    session.clear()
    flash(f"Backup restored ({restored.get('delivery_notes',0)} delivery notes, {restored.get('companies',0)} companies, {nf} image files). Please log in again.", 'success')
    return redirect(url_for('login'))


# ── Auto-initialize DB on startup (works with gunicorn + Railway) ──────────
import logging as _logging
_logging.basicConfig(level=_logging.INFO)
_logging.getLogger(__name__).info("Running init_db() at startup...")
init_db()

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print("=" * 55)
    print("  Geometry Home Delivery Management System")
    print(f"  Running at: http://localhost:{port}")
    print("  Login: admin / admin123")
    print("=" * 55)
    app.run(debug=False, host='0.0.0.0', port=port)
